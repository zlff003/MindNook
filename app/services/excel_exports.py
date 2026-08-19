from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="173F3A")
TITLE_FILL = PatternFill("solid", fgColor="E8F3F0")
RISK_FILL = PatternFill("solid", fgColor="FDE8E7")
CONSULT_FILL = PatternFill("solid", fgColor="FFF4D6")
THIN_BORDER = Border(bottom=Side(style="thin", color="D7E3E0"))


def _safe_cell(value):
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def build_excel_export(
    *,
    title: str,
    sheet_name: str,
    headers: Sequence[str],
    rows: Iterable[Sequence],
    notice: str = "请妥善保管导出文件并控制访问范围。",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]
    sheet.sheet_view.showGridLines = False

    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="173F3A")
    sheet["A1"].fill = TITLE_FILL
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 32

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = (
        f"导出时间：{datetime.now(UTC).strftime('%Y-%m-%d %H:%M:%S UTC')}；{notice}"
    )
    sheet["A2"].font = Font(name="Microsoft YaHei", size=9, color="5C6F6B")
    sheet["A2"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 24

    header_row = 4
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = HEADER_FILL
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[header_row].height = 24

    materialized_rows = list(rows)
    for row_index, row in enumerate(materialized_rows, start=header_row + 1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column, _safe_cell(value))
            cell.font = Font(name="Microsoft YaHei", size=10, color="243532")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if value == "RISK":
                cell.fill = RISK_FILL
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="A62925")
            elif value == "CONSULT":
                cell.fill = CONSULT_FILL

    data_end = max(header_row, header_row + len(materialized_rows))
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{data_end}"
    sheet.freeze_panes = f"A{header_row + 1}"

    for column in range(1, len(headers) + 1):
        values = [str(headers[column - 1])]
        values.extend(
            str(row[column - 1])
            for row in materialized_rows
            if column - 1 < len(row) and row[column - 1] is not None
        )
        width = min(max(max((len(item) for item in values), default=8) + 2, 12), 42)
        sheet.column_dimensions[get_column_letter(column)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
