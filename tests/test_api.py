import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.agents.runtime import AgentRuntime
from app.core.config import Settings
from app.main import create_app
from app.models import (
    AlertDelivery,
    AgentCollaborationEvent,
    AgentRun,
    AgentRunTrace,
    AgentStep,
    AuditLog,
    CaseEvent,
    LedgerEntry,
    Message,
    PsychologicalReport,
    RiskCase,
    ToolJob,
    User,
)
from app.services.model import MockModelProvider


class RecordingModelProvider(MockModelProvider):
    def __init__(self):
        self.seen_contents: list[str] = []

    def complete(self, messages, *, model, temperature, json_mode, purpose):
        self.seen_contents.extend(message.content for message in messages)
        return super().complete(
            messages,
            model=model,
            temperature=temperature,
            json_mode=json_mode,
            purpose=purpose,
        )


class MindNookApiTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        database_path = Path(self.temp_dir.name) / "mindnook-test.db"
        settings = Settings(
            app_env="test",
            app_secret="test-secret-that-is-not-used-in-production",
            database_url=f"sqlite:///{database_path.as_posix()}",
            model_provider="mock",
            demo_student_username="student",
            demo_student_password="student123",
            demo_admin_username="admin",
            demo_admin_password="admin123",
            smtp_host="",
            smtp_username="",
            smtp_password="",
            smtp_from_address="",
            alert_email_to="",
            smtp_use_tls=False,
            smtp_use_ssl=False,
        )
        self.client_context = TestClient(create_app(settings))
        self.client = self.client_context.__enter__()
        self.student_token = self._login("student", "student123")
        self.admin_token = self._login("admin", "admin123")

    def tearDown(self):
        self.client_context.__exit__(None, None, None)
        self.temp_dir.cleanup()

    def _login(self, username: str, password: str) -> str:
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200)
        return response.json()["access_token"]

    def _headers(self, token: str) -> dict[str, str]:
        return {"Authorization": f"Bearer {token}"}

    def _create_conversation(self) -> str:
        response = self.client.post(
            "/api/v1/conversations",
            headers=self._headers(self.student_token),
        )
        self.assertEqual(response.status_code, 200)
        return response.json()["id"]

    def _chat(self, conversation_id: str, content: str, request_id: str):
        return self.client.post(
            f"/api/v1/conversations/{conversation_id}/messages:stream",
            headers=self._headers(self.student_token),
            json={
                "content": content,
                "client_request_id": request_id,
            },
        )

    def _count(self, model) -> int:
        session = self.client.app.state.session_factory()
        try:
            return session.scalar(select(func.count()).select_from(model)) or 0
        finally:
            session.close()

    def test_pages_and_health_are_available(self):
        self.assertEqual(self.client.get("/health/live").status_code, 200)
        student_page = self.client.get("/student")
        self.assertEqual(student_page.status_code, 200)
        self.assertIn('id="historyToggleButton"', student_page.text)
        self.assertIn('id="stopButton"', student_page.text)
        self.assertIn('id="privacyDialog"', student_page.text)
        admin_page = self.client.get("/admin")
        self.assertEqual(admin_page.status_code, 200)
        for route in ("overview", "reports", "cases", "alerts", "operations", "audit", "users"):
            self.assertIn(f'data-route="{route}"', admin_page.text)
            self.assertIn(f'data-page="{route}"', admin_page.text)
        self.assertIn('id="detailDrawer"', admin_page.text)
        self.assertIn('aria-modal="true"', admin_page.text)
        self.assertIn('id="reportPagination"', admin_page.text)
        self.assertIn('id="exportReportsButton"', admin_page.text)
        self.assertIn('id="exportCasesButton"', admin_page.text)
        self.assertIn('id="operationOverviewPanel"', admin_page.text)
        self.assertIn('id="jobGroupList"', admin_page.text)
        self.assertIn('id="exportAuditButton"', admin_page.text)
        admin_script = self.client.get("/static/admin/admin.js")
        self.assertEqual(admin_script.status_code, 200)
        self.assertIn('EXCEL_REPORT: "excel_report"', admin_script.text)
        self.assertIn('excel_report: "写入心理报告"', admin_script.text)
        self.assertIn('return "心理报告写入流程"', admin_script.text)

    def test_student_registration_and_admin_created_account_use_institutional_ids(self):
        student_response = self.client.post(
            "/api/v1/auth/register/student",
            json={
                "institutional_id": "20260001",
                "name": "张同学",
                "password": "student-pass-2026",
            },
        )
        self.assertEqual(student_response.status_code, 201)
        self.assertEqual(student_response.json()["user"]["institutional_id"], "20260001")
        self.assertEqual(student_response.json()["user"]["name"], "张同学")
        self.assertNotIn("password", student_response.text)
        duplicate_response = self.client.post(
            "/api/v1/auth/register/student",
            json={
                "institutional_id": "20260001",
                "name": "另一位同学",
                "password": "student-pass-2026",
            },
        )
        self.assertEqual(duplicate_response.status_code, 409)

        forbidden_response = self.client.post(
            "/api/v1/admin/users/admins",
            headers=self._headers(student_response.json()["access_token"]),
            json={
                "institutional_id": "T9001",
                "name": "未授权管理员",
                "password": "admin-pass-2026",
            },
        )
        self.assertEqual(forbidden_response.status_code, 403)

        admin_response = self.client.post(
            "/api/v1/admin/users/admins",
            headers=self._headers(self.admin_token),
            json={
                "institutional_id": "T9001",
                "name": "王老师",
                "password": "admin-pass-2026",
            },
        )
        self.assertEqual(admin_response.status_code, 201)
        login_response = self.client.post(
            "/api/v1/auth/login",
            json={"institutional_id": "T9001", "password": "admin-pass-2026"},
        )
        self.assertEqual(login_response.status_code, 200)
        self.assertEqual(login_response.json()["user"]["role"], "admin")
        session = self.client.app.state.session_factory()
        try:
            registered = session.scalar(
                select(User).where(User.institutional_id == "20260001")
            )
            self.assertEqual(registered.name, "张同学")
            self.assertNotEqual(registered.password_hash, "student-pass-2026")
        finally:
            session.close()

    def test_only_two_roles_can_log_in_to_their_apis(self):
        student_admin_response = self.client.get(
            "/api/v1/admin/reports",
            headers=self._headers(self.student_token),
        )
        self.assertEqual(student_admin_response.status_code, 403)

        admin_student_response = self.client.post(
            "/api/v1/conversations",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(admin_student_response.status_code, 403)

    def test_unauthenticated_requests_cannot_access_student_or_admin_apis(self):
        self.assertEqual(self.client.get("/api/v1/conversations").status_code, 401)
        self.assertEqual(self.client.get("/api/v1/admin/reports").status_code, 401)

    def test_students_cannot_read_each_others_conversations(self):
        conversation_id = self._create_conversation()
        self._chat(
            conversation_id,
            "最近考试压力很大",
            "request-owner-isolation-0001",
        )
        second_student = self.client.post(
            "/api/v1/auth/register/student",
            json={
                "institutional_id": "20269999",
                "name": "隔离测试学生",
                "password": "student-pass-2026",
            },
        )
        self.assertEqual(second_student.status_code, 201)

        response = self.client.get(
            f"/api/v1/conversations/{conversation_id}/messages",
            headers=self._headers(second_student.json()["access_token"]),
        )

        self.assertEqual(response.status_code, 404)
        self.assertNotIn("最近考试压力很大", response.text)

    def test_general_chat_streams_without_report(self):
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "请帮我解释 Python 循环",
            "request-general-0001",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("event: accepted", response.text)
        self.assertIn("event: delta", response.text)
        self.assertIn("event: completed", response.text)
        self.assertEqual(self._count(PsychologicalReport), 0)
        self.assertEqual(self._count(RiskCase), 0)
        self.assertEqual(self._count(AlertDelivery), 0)

    def test_consult_creates_report_and_ledger_only(self):
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "最近考试压力很大，晚上睡不着",
            "request-support-0001",
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("风险等级", response.text)
        self.assertEqual(self._count(PsychologicalReport), 1)
        self.assertEqual(self._count(LedgerEntry), 1)
        self.assertEqual(self._count(RiskCase), 0)
        self.assertEqual(self._count(AlertDelivery), 0)

    def test_persistent_impact_remains_consult_without_case(self):
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "我每天都很焦虑，已经无法学习",
            "request-medium-0001",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._count(PsychologicalReport), 1)
        self.assertEqual(self._count(LedgerEntry), 1)
        self.assertEqual(self._count(RiskCase), 0)
        self.assertEqual(self._count(AlertDelivery), 0)

    def test_risk_creates_full_follow_up_chain(self):
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "我想自杀，现在完全撑不住了",
            "request-high-0001",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("不要独处", response.text)
        self.assertIn("信任的人", response.text)
        self.assertNotIn("你的风险等级", response.text)
        self.assertEqual(self._count(PsychologicalReport), 1)
        self.assertEqual(self._count(LedgerEntry), 1)
        self.assertEqual(self._count(RiskCase), 1)
        self.assertEqual(self._count(AlertDelivery), 1)

    def test_admin_groups_terminal_jobs_hides_active_work_and_retries_failures(self):
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "我想自杀，现在完全撑不住了",
            "request-job-group-0001",
        )
        self.assertEqual(response.status_code, 200)

        groups = self.client.get(
            "/api/v1/admin/job-groups",
            headers=self._headers(self.admin_token),
        ).json()
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["status"], "SUCCESS")
        self.assertEqual(
            [step["job_type"] for step in groups[0]["steps"]],
            ["CASE_CREATE", "ALERT_SEND"],
        )

        session = self.client.app.state.session_factory()
        try:
            alert_job = session.scalar(
                select(ToolJob).where(ToolJob.job_type == "ALERT_SEND")
            )
            alert_job.status = "PENDING"
            session.commit()
        finally:
            session.close()
        hidden_groups = self.client.get(
            "/api/v1/admin/job-groups",
            headers=self._headers(self.admin_token),
        ).json()
        self.assertEqual(hidden_groups, [])

        session = self.client.app.state.session_factory()
        try:
            alert_job = session.scalar(
                select(ToolJob).where(ToolJob.job_type == "ALERT_SEND")
            )
            alert_job.status = "DEAD"
            alert_job.attempt_count = alert_job.max_attempts
            alert_job.last_error = "smtp connection timed out"
            session.commit()
        finally:
            session.close()
        failed_groups = self.client.get(
            "/api/v1/admin/job-groups",
            headers=self._headers(self.admin_token),
        ).json()
        self.assertEqual(len(failed_groups), 1)
        self.assertEqual(failed_groups[0]["status"], "DEAD")
        self.assertEqual(failed_groups[0]["steps"][0]["status"], "SUCCESS")
        self.assertEqual(failed_groups[0]["steps"][1]["status"], "DEAD")

        retry = self.client.post(
            f"/api/v1/admin/job-groups/{failed_groups[0]['report_id']}/retry",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(len(retry.json()["retried_jobs"]), 1)

        completed_groups = self.client.get(
            "/api/v1/admin/job-groups",
            headers=self._headers(self.admin_token),
        ).json()
        self.assertEqual(completed_groups[0]["status"], "SUCCESS")
        session = self.client.app.state.session_factory()
        try:
            case_job = session.scalar(
                select(ToolJob).where(ToolJob.job_type == "CASE_CREATE")
            )
            alert_job = session.scalar(
                select(ToolJob).where(ToolJob.job_type == "ALERT_SEND")
            )
            retry_audit = session.scalar(
                select(AuditLog).where(AuditLog.action == "job_group_retried")
            )
            self.assertEqual(case_job.attempt_count, 1)
            self.assertEqual(alert_job.attempt_count, 1)
            self.assertIsNotNone(retry_audit)
        finally:
            session.close()

        repeated_retry = self.client.post(
            f"/api/v1/admin/job-groups/{failed_groups[0]['report_id']}/retry",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(repeated_retry.status_code, 409)

    def test_admin_retry_restores_a_fully_failed_dependent_job_group(self):
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "我想自杀，现在完全撑不住了",
            "request-job-group-0002",
        )
        self.assertEqual(response.status_code, 200)

        session = self.client.app.state.session_factory()
        try:
            jobs = list(session.scalars(select(ToolJob).order_by(ToolJob.id.asc())))
            self.assertEqual(len(jobs), 2)
            self.assertEqual(jobs[1].depends_on_job_id, jobs[0].id)
            for job in jobs:
                job.status = "DEAD"
                job.attempt_count = job.max_attempts
                job.last_error = f"forced {job.job_type} failure"
            report_id = session.get(PsychologicalReport, jobs[0].report_id).public_id
            session.commit()
        finally:
            session.close()

        retry = self.client.post(
            f"/api/v1/admin/job-groups/{report_id}/retry",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(len(retry.json()["retried_jobs"]), 2)

        groups = self.client.get(
            "/api/v1/admin/job-groups",
            headers=self._headers(self.admin_token),
        ).json()
        self.assertEqual(groups[0]["status"], "SUCCESS")
        self.assertEqual(
            [step["status"] for step in groups[0]["steps"]],
            ["SUCCESS", "SUCCESS"],
        )
        self.assertEqual(self._count(RiskCase), 1)
        self.assertEqual(self._count(AlertDelivery), 1)

        session = self.client.app.state.session_factory()
        try:
            audit = session.scalar(
                select(AuditLog).where(AuditLog.action == "job_group_retried")
            )
            self.assertIn("retried_jobs=2", audit.detail)
        finally:
            session.close()

    def test_duplicate_request_is_idempotent(self):
        conversation_id = self._create_conversation()
        first = self._chat(
            conversation_id,
            "最近考试压力很大",
            "request-duplicate-0001",
        )
        second = self._chat(
            conversation_id,
            "最近考试压力很大",
            "request-duplicate-0001",
        )
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertIn('"duplicate": true', second.text)
        self.assertEqual(self._count(PsychologicalReport), 1)
        self.assertEqual(self._count(LedgerEntry), 1)
        self.assertEqual(self._count(Message), 2)

    def test_report_context_access_is_audited(self):
        conversation_id = self._create_conversation()
        self._chat(
            conversation_id,
            "最近考试压力很大",
            "request-audit-0001",
        )
        reports = self.client.get(
            "/api/v1/admin/reports",
            headers=self._headers(self.admin_token),
        ).json()
        response = self.client.get(
            f"/api/v1/admin/reports/{reports[0]['id']}/context",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._count(AuditLog), 1)

    def test_agent_run_and_steps_are_persisted_and_visible_to_admin(self):
        conversation_id = self._create_conversation()
        self._chat(
            conversation_id,
            "最近考试压力很大，晚上睡不着",
            "request-agent-trace-0001",
        )

        self.assertEqual(self._count(AgentRun), 1)
        self.assertEqual(self._count(AgentRunTrace), 1)
        self.assertGreaterEqual(self._count(AgentStep), 6)
        self.assertGreater(self._count(AgentCollaborationEvent), 20)
        runs_response = self.client.get(
            "/api/v1/admin/agent-runs",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(runs_response.status_code, 200)
        run = runs_response.json()[0]
        self.assertEqual(run["risk_level"], "CONSULT")
        self.assertEqual(run["provider"], "mock")
        self.assertFalse(run["degraded"])

        steps_response = self.client.get(
            f"/api/v1/admin/agent-runs/{run['id']}/steps",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(steps_response.status_code, 200)
        step_names = [item["step_name"] for item in steps_response.json()]
        self.assertIn("hard_safety_rules", step_names)
        self.assertIn("response_review", step_names)

        events_response = self.client.get(
            f"/api/v1/admin/agent-runs/{run['id']}/events",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(events_response.status_code, 200)
        event_types = [item["event_type"] for item in events_response.json()]
        self.assertIn("TASK_CLAIMED", event_types)
        self.assertIn("ARTIFACT_PUBLISHED", event_types)
        self.assertIn("FINAL_ADOPTED", event_types)

        trace_response = self.client.get(
            f"/api/v1/admin/agent-runs/{run['id']}/trace",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(trace_response.status_code, 200)
        audit_memory = trace_response.json()
        self.assertTrue(audit_memory["trace_available"])
        self.assertIsNone(audit_memory["trace_notice"])
        self.assertIn("memoryBrief", audit_memory)
        self.assertIn("modelHistory", audit_memory)
        self.assertTrue(audit_memory["agent_steps"])
        self.assertTrue(audit_memory["collaboration_events"])
        self.assertTrue(audit_memory["tasks"])
        self.assertTrue(audit_memory["artifacts"])
        risk_decisions = audit_memory["risk_decisions"]
        self.assertEqual(
            [item["stage"] for item in risk_decisions],
            ["understanding", "hard_rule", "safety", "run_final"],
        )
        understanding = risk_decisions[0]
        self.assertEqual(understanding["producer"], "UnderstandingAgent")
        self.assertEqual(understanding["risk_level"], "CONSULT")
        self.assertIsNotNone(understanding["confidence"])
        self.assertTrue(understanding["reason"])
        self.assertEqual(risk_decisions[-1]["risk_level"], run["risk_level"])

        session = self.client.app.state.session_factory()
        try:
            persisted_trace = session.scalar(
                select(AgentRunTrace).where(AgentRunTrace.run_id == session.scalar(
                    select(AgentRun.id).where(AgentRun.public_id == run["id"])
                ))
            )
            session.delete(persisted_trace)
            session.commit()
        finally:
            session.close()
        legacy_response = self.client.get(
            f"/api/v1/admin/agent-runs/{run['id']}/trace",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(legacy_response.status_code, 200)
        legacy_trace = legacy_response.json()
        self.assertFalse(legacy_trace["trace_available"])
        self.assertIn("无法恢复 UnderstandingAgent", legacy_trace["trace_notice"])
        self.assertEqual(legacy_trace["risk_decisions"], [])

    def test_raw_pii_never_enters_agent_model_messages(self):
        provider = RecordingModelProvider()
        runtime = AgentRuntime(
            provider,
            self.client.app.state.settings,
        )
        self.client.app.state.agent_runtime = runtime
        self.client.app.state.agent_harness.runtime = runtime
        conversation_id = self._create_conversation()
        response = self._chat(
            conversation_id,
            "我的电话是13800138000，最近考试压力很大",
            "request-redaction-0001",
        )

        self.assertEqual(response.status_code, 200)
        self.assertNotIn("13800138000", "\n".join(provider.seen_contents))
        self.assertIn("[PHONE]", "\n".join(provider.seen_contents))

        run = self.client.get(
            "/api/v1/admin/agent-runs",
            headers=self._headers(self.admin_token),
        ).json()[0]
        trace = self.client.get(
            f"/api/v1/admin/agent-runs/{run['id']}/trace",
            headers=self._headers(self.admin_token),
        ).json()
        self.assertIn("13800138000", trace["source_input"])
        self.assertNotIn("13800138000", trace["redacted_input"])
        self.assertIn("[PHONE]", trace["redacted_input"])

    def test_case_notes_and_status_keep_append_only_history(self):
        conversation_id = self._create_conversation()
        self._chat(
            conversation_id,
            "我想自杀，现在完全撑不住了",
            "request-case-history-0001",
        )
        cases = self.client.get(
            "/api/v1/admin/cases",
            headers=self._headers(self.admin_token),
        ).json()
        case_id = cases[0]["id"]

        note_response = self.client.post(
            f"/api/v1/admin/cases/{case_id}/notes",
            headers=self._headers(self.admin_token),
            json={"note": "已联系学生，安排进一步沟通。"},
        )
        transition_response = self.client.post(
            f"/api/v1/admin/cases/{case_id}/transitions",
            headers=self._headers(self.admin_token),
            json={"status": "following"},
        )
        history_response = self.client.get(
            f"/api/v1/admin/cases/{case_id}/events",
            headers=self._headers(self.admin_token),
        )

        self.assertEqual(note_response.status_code, 200)
        self.assertEqual(transition_response.status_code, 200)
        self.assertEqual(history_response.status_code, 200)
        event_types = [item["event_type"] for item in history_response.json()]
        self.assertEqual(event_types, ["created", "note_added", "status_changed"])
        self.assertEqual(self._count(CaseEvent), 3)

    def test_case_owner_audit_center_and_excel_exports(self):
        conversation_id = self._create_conversation()
        source_text = "我想自杀，现在完全撑不住了\n我的手机号码是13800138000"
        self._chat(conversation_id, source_text, "request-owner-export-0001")
        owner = self.client.post(
            "/api/v1/admin/users/admins",
            headers=self._headers(self.admin_token),
            json={
                "institutional_id": "T8001",
                "name": "李老师",
                "password": "owner-pass-2026",
            },
        ).json()
        case_item = self.client.get(
            "/api/v1/admin/cases",
            headers=self._headers(self.admin_token),
        ).json()[0]
        assignment = self.client.post(
            f"/api/v1/admin/cases/{case_item['id']}/assign",
            headers=self._headers(self.admin_token),
            json={"owner_id": owner["id"]},
        )
        self.assertEqual(assignment.status_code, 200)
        self.assertEqual(assignment.json()["owner"]["institutional_id"], "T8001")
        self.assertEqual(assignment.json()["status"], "assigned")

        history = self.client.get(
            f"/api/v1/admin/cases/{case_item['id']}/events",
            headers=self._headers(self.admin_token),
        ).json()
        assignment_event = next(
            event for event in history if event["event_type"] == "owner_assigned"
        )
        self.assertEqual(assignment_event["actor"]["institutional_id"], "admin")

        audit_response = self.client.get(
            "/api/v1/admin/audit-logs?page=1&page_size=100&resource_type=case",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(audit_response.status_code, 200)
        audit_count = self._count(AuditLog)
        audit_items = audit_response.json()["items"]
        self.assertIn(
            "case_owner_assigned",
            [item["action"] for item in audit_items],
        )
        assignment_audit = next(
            item for item in audit_items if item["action"] == "case_owner_assigned"
        )
        self.assertEqual(assignment_audit["action_label"], "分配或转交个案")
        self.assertEqual(assignment_audit["resource_type_label"], "风险个案")
        self.assertIn("负责人：T8001", assignment_audit["detail_label"])

        second_audit_response = self.client.get(
            "/api/v1/admin/audit-logs?page=1&page_size=100",
            headers=self._headers(self.admin_token),
        )
        self.assertEqual(second_audit_response.status_code, 200)
        self.assertEqual(self._count(AuditLog), audit_count)

        export_values = {}
        for export_type in ("reports", "cases", "audit-logs"):
            response = self.client.get(
                f"/api/v1/admin/exports/{export_type}.xlsx",
                headers=self._headers(self.admin_token),
            )
            self.assertEqual(response.status_code, 200)
            self.assertIn("spreadsheetml", response.headers["content-type"])
            workbook = load_workbook(BytesIO(response.content), read_only=True)
            export_values[export_type] = [
                str(value)
                for row in workbook.active.iter_rows(values_only=True)
                for value in row
                if value is not None
            ]

        for export_type in ("reports", "cases"):
            values = export_values[export_type]
            self.assertIn("学生原始表达", values)
            self.assertIn(source_text, values)
            self.assertIn(
                "本文件完整包含学生原始表达等敏感原文，请妥善保管并控制访问范围。",
                values[1],
            )
        self.assertTrue(
            any("原始表达摘录：" in value for value in export_values["cases"])
        )
        self.assertIn("分配或转交个案", export_values["audit-logs"])
        self.assertIn("风险个案", export_values["audit-logs"])


if __name__ == "__main__":
    unittest.main()
