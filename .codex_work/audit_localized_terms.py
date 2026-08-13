from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document


ROOT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\localized_route_terms")
OLD = re.compile(r"\b(?:CHAT|CONSULT|RISK)\b")
NEW = ("日常交流", "心理支持", "风险处置")
SUSPECT = re.compile(
    r"心理支持支持|日常交流交流|风险处置风险|风险处置处置|"
    r"风险处置\s+(?:流程|生成|信息|闭环|场景|保存|回复|的|与|对话|候选|优先)|"
    r"心理支持\s+(?:支持|场景|生成|保存|回复|的|只)|"
    r"日常交流\s+(?:无|不得|不)|真实风险处置|风险处置(?:召回率|错误降级率|夜间值守|表达|固定规则)"
)


def iter_text(doc):
    for i, p in enumerate(doc.paragraphs):
        yield f"P{i:04d}", p.text.strip()
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            yield f"T{ti}R{ri}", " | ".join(c.text.strip().replace("\n", " / ") for c in row.cells)


for path in sorted(ROOT.glob("*.docx")):
    doc = Document(path)
    lines = [(loc, text) for loc, text in iter_text(doc) if text]
    old_hits = [(loc, text) for loc, text in lines if OLD.search(text)]
    suspect_hits = [(loc, text) for loc, text in lines if SUSPECT.search(text)]
    relevant = [(loc, text) for loc, text in lines if any(term in text for term in NEW)]
    with ZipFile(path) as archive:
        broken = archive.testzip()
        old_parts = []
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                value = archive.read(name).decode("utf-8", "ignore")
                if OLD.search(value):
                    old_parts.append(name)
    print(f"\n===== {path.name} =====")
    print("STRUCTURE", {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "zip_error": broken})
    print("OLD", len(old_hits), "OLD_XML", old_parts, "SUSPECT", len(suspect_hits), "RELEVANT", len(relevant))
    for loc, text in suspect_hits:
        print("SUSPECT", loc, text, sep="\t")
    for loc, text in relevant:
        print("TERM", loc, text, sep="\t")
