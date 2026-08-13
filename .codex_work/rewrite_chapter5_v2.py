from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\Projects\Agent\MindBridge\项目提交文档\MindNook智能体 - 设计方案书 - 第5章更新版.docx")
OUTPUT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\MindNook智能体 - 设计方案书 - 第5章重构版.docx")


def find_paragraph(document, exact_text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact_text}")


def insert_before(anchor, text, style):
    node = OxmlElement("w:p")
    anchor._p.addprevious(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def delete_range(document, start_text, end_text):
    start = find_paragraph(document, start_text)
    end = find_paragraph(document, end_text)
    node = start._p
    while node is not None and node is not end._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node


def main():
    document = Document(SOURCE)
    delete_range(
        document,
        "5 模型、上下文、知识与 Skill 设计",
        "6 安全、工具与业务闭环",
    )
    anchor = find_paragraph(document, "6 安全、工具与业务闭环")

    content = [
        ("Heading 1", "5 模型、上下文与知识设计"),
        ("Normal", "本章说明 MindNook 如何把模型推理、Prompt 约束、会话上下文、分层记忆、知识检索与心理场景 Skill 组织为统一的回复支撑体系。整体设计遵循“模型负责理解与生成、上下文提供必要连续性、RAG 提供可追溯事实、Skill 提供场景表达规则、安全机制设置不可绕过的边界”的原则。各类信息由 ContextAgent 汇总为结构化 ContextBundle，再由 ResponseAgent 生成候选回复，并由 SafetyAgent 在公开前完成审查。模型、知识或 Skill 的任何输出都不能自行改变风险等级，也不能绕过固定安全规则和 CoordinatorAgent 的最终采纳控制。"),

        ("Heading 2", "5.1 Prompt 与结构化输出"),
        ("Normal", "MindNook 的 Prompt 不采用将角色设定、历史消息、知识全文和场景规则一次性拼接的方式，而是根据 Agent 职责和当前任务进行分层装配。第一层是不可覆盖的通用约束，包括角色职责、输入输出边界、隐私要求、禁止临床诊断、禁止药物建议、禁止危险操作细节、禁止泄露后台风险信息以及禁止编造校园资源等规则；第二层是本轮运行状态，包括脱敏后的当前输入、intent、risk、memoryBrief、modelHistory 和必要的协作产物；第三层是 ContextAgent 按风险等级和主题选出的 RAG 证据与 Skill Workflow；第四层是当前任务指令，例如识别意图、评估风险、生成候选回复或审查指定版本。后加载的知识和 Skill 只能补充内容，不能覆盖上层安全约束。"),
        ("Normal", "不同智能体使用相互独立的 Prompt 模板。UnderstandingAgent 只识别对话类型、主题、知识需求、置信度和判断理由，不生成安慰或处置建议；SafetyAgent 在风险评估阶段关注危险表达、即时危险和安全问题，在回复审查阶段检查候选文本是否违反固定规则；ContextAgent 负责选择并组织上下文，不直接形成面向学生的结论；ResponseAgent 只依据已发布的 hard_rule、intent、risk 和 ContextBundle 生成候选回复。职责隔离能够减少一个模型调用同时承担“理解、判断、检索、生成、审查”所造成的目标冲突，也使每类输出可以单独测试和追踪。"),
        ("Normal", "模型间交换的数据采用结构化输出。意图产物保存对话类型、主题、知识需求、置信度和判断理由；风险产物保存风险等级、情绪概况、即时危险、是否需要安全确认及判定依据；ContextBundle 保存 memoryBrief、modelHistory、检索证据、已加载 Skill、回复约束和检索降级标记；回复审查结果保存 approved、violations、revision_guidance 以及被审查的候选版本。结构化字段进入 AgentArtifact 后只追加新版本，不覆盖旧结果，从而避免自由文本被后续组件误读，并保证审查结论与具体 response_proposal 一一对应。"),
        ("Normal", "ResponseAgent 的正文首先作为内部 response_proposal 发布，不能直接发送给学生。SafetyAgent 必须同时执行模型审查和确定性规则检查，重点识别诊断或药物建议、危险细节、虚构资源、绝对保密承诺、歧视标签、后台信息泄露，以及风险场景中现实安全行动不足或问题过多等情况。只有审查结果针对当前候选版本且 approved=true 时，CoordinatorAgent 才能发布 final_response；首次不通过时根据 violations 和 revision_guidance 进行一次受限修复，再次不通过或预算耗尽则采用固定安全回退。"),
        ("Normal", "Prompt 装配坚持最小必要原则。手机号、邮箱、身份证号、学号等敏感信息在进入模型和短期记忆前进行脱敏；工作人员交接摘要、数据库实体、工具任务状态和其他 Agent 的私有记忆不进入学生端回复 Prompt。RAG 文本只作为事实参考，不作为系统指令执行；未加载 Skill 的规则不得由模型自行补写；知识库未明确提供的人员、电话、办公时间和政策不得推断生成。该方式既控制 Token 规模，也降低提示注入、后台信息泄露和历史判断固化的风险。"),

        ("Heading 2", "5.2 模型配置与调用"),
        ("Normal", "模型能力通过统一的 ModelProvider 接口接入，当前支持 Mock 模式以及基于 OpenAI-compatible API 的模型服务配置。业务代码只依赖统一的消息、结构化输出和流式生成接口，不直接依赖特定厂商 SDK，因此可以在测试环境使用确定性的 MockModelProvider，在部署环境切换为实际模型服务，而无需改变 AgentTask、AgentArtifact、ContextBundle 和安全审查协议。模型标识、访问地址、鉴权信息、超时和采样参数均由配置注入，不写入 Prompt、数据库正文或前端代码。"),
        ("Normal", "UnderstandingAgent、SafetyAgent 和 ResponseAgent 可以按职责分别配置模型与参数。意图识别、风险评估和回复审查属于结构化判断任务，优先采用较低随机性，保证字段稳定、便于回归；回复生成需要兼顾自然度和一致性，可以保留适度语言变化，但仍受篇幅、建议数量、提问数量和安全规则约束。ContextAgent 以确定性程序装配记忆、RAG 与 Skill，避免再增加一次没有必要的模型调用。通过按任务拆分模型配置，可以在成本、时延、表达质量与安全稳定性之间进行独立调整。"),
        ("Normal", "CoordinatorAgent 对单轮模型调用实行统一预算控制。默认最大协作轮数为 16，单轮最多调用模型 7 次，模型单次超时为 20 秒，整轮模型运行时限为 55 秒。每次调用前检查剩余预算，调用后记录 Agent、模型、耗时、Token 估算、成功状态和错误摘要；达到轮数、次数或时间上限后不继续开放式循环，而是使用当前已有 Artifact 或固定安全模板结束本轮。回复修复最多执行一次，防止“生成—拒绝—重写”形成无界循环。"),
        ("Normal", "回复流采用“内部生成、审查通过后公开”的方式。ResponseAgent 可以流式产生 Token，但这些 Token 只在服务端累积为候选文本；SafetyAgent 完成整段审查并由 CoordinatorAgent 采纳后，AgentHarness 才通过公开 SSE 事件按片段返回最终成品。这样既保留流式交互体验，也避免不安全内容在审查完成前先行流出。模型调用日志和内部候选只进入 Trace 与审计链路，不作为学生可见事件。"),
        ("Normal", "模型失败时按照任务类型确定性降级。意图解析失败时由固定规则形成降级 intent；模型风险评估失败时仍保留固定安全规则给出的风险下限，明确危险表达不会因模型不可用而降级；回复生成失败时使用带 trusted_fallback 标记的固定模板；心理支持或风险处置的安全审查服务不可用时进入固定安全回退，只有日常交流候选通过确定性输出规则后才允许放行。所有超时、解析失败、预算耗尽和降级结果均记录在 AgentStep、Event 与 RuntimeTrace 中，不将故障伪装为正常模型结论。"),

        ("Heading 2", "5.3 上下文与分层记忆"),
        ("Normal", "AgentHarness 在调用 Runtime 前构造 MemoryContext，并把当前会话的可信历史与本轮输入分开装配。MySQL 中的 chat_messages 保存完整业务消息，是会话记录的事实来源；Redis 使用 mindbridge:short-term-memory:<session_public_id> 保存当前会话最近最多 40 条脱敏消息，默认 TTL 为 86400 秒。Redis 命中时直接读取并刷新 TTL；Key 缺失、内容损坏或服务不可用时，从 MySQL 回填当前会话最近 40 条 completed 消息，恢复时间顺序后继续运行。不同学生和不同会话始终使用独立 Key，禁止使用其他会话内容补全上下文。"),
        ("Normal", "长对话采用“历史摘要 + 最近原文窗口”的组合策略。较早历史通过确定性压缩器形成最多 500 字符的 memoryBrief，固定保留最多 4 条学生关注、3 条助手支持和本轮输入三个栏目；最近最多 8 条带角色的脱敏原文组成 modelHistory，并受 6000 字符总预算限制。短历史只注入角色化原文，避免摘要与原文重复；历史超过窗口时，memoryBrief 作为独立 system 背景加入 Prompt，帮助模型保持主题连续性，同时把注意力留给当前输入。"),
        ("Normal", "摘要内容只描述学生已经表达的关注、系统已经提供的支持以及本轮需要关注的事项，不生成诊断标签，不推断未表达的经历，也不把后台风险等级写成长期画像。单条历史过长时保留开头与结尾并省略中段；整体超出预算时优先保留栏目结构和较新的信息。数据库中的完整消息不会因压缩被删除，但模型本轮可见范围仍限制在最近 40 条之内，更早历史不会进入 memoryBrief，从数据最小化角度控制长期积累对当前回复的影响。"),
        ("Normal", "除共享会话记忆外，各专业 Agent 还可以使用角色私有记忆。私有记忆以 agent:<AgentName>:<session_public_id> 隔离保存，默认每个 Key 最多 20 条、TTL 为 86400 秒，构造单个 Agent 的 Prompt 时只读取该角色最近 4 条记录。UnderstandingAgent 的历史意图、SafetyAgent 的风险判断、ContextAgent 的检索策略和 ResponseAgent 的表达策略仅供本角色参考，不允许跨角色直接注入。私有记忆不能覆盖当前输入和固定规则，也不能成为对学生的持久化风险标签。"),
        ("Normal", "ContextAgent 在完成记忆装配后，将 memoryBrief、modelHistory、知识证据、Skill Workflow 和回复约束统一写入 ContextBundle，并发布 context_ready，供 ResponseAgent 使用。完整的 Agent Run、协作步骤、黑板事件、检索阶段和模型调用信息进入 RuntimeTrace，用于审计与问题定位，但 Trace 永不回流下一轮 Prompt。最终数据库事务提交成功后，学生消息与助手成品才追加到 Redis 并裁剪至 40 条；若事务回滚，则不把未完成回复写入短期记忆，保证业务记录、模型上下文和学生实际收到的内容一致。"),

        ("Heading 2", "5.4 混合 RAG 检索"),
        ("Normal", "心理知识检索由 ContextAgent 按最终风险等级确定是否触发。日常交流不调用心理 RAG，避免把学习、编程、校园事务等普通问题过度心理化；心理支持根据学生当前困扰、UnderstandingAgent 提取的主题和知识需求构造 Query；风险处置除当前输入和主题外，固定补充“当前安全、危机支持、现实求助”等安全主题，使检索结果优先服务于短时间内的现实安全行动。RAG 只提供知识证据，不参与最终风险定级，也不能降低固定规则确定的安全下限。"),
        ("Normal", "知识源采用 UTF-8 Markdown，入库时保留文件名、章节路径和块序号等元数据。文本按章节和自然语言边界切分，默认 Chunk 大小为 512 字符、重叠 64 字符；切分时优先在段落、句号、问号、感叹号和分号附近结束，避免破坏完整语义。知识重建任务同步生成 Chroma 向量索引和 BM25 词法索引，使语义相近但词面不同的表达可以被向量召回，校园术语、专有名词和明确关键词可以被 BM25 稳定命中。"),
        ("Normal", "在线检索同时执行向量召回和 BM25 召回，两路各最多取得 candidate_k=12 个候选，并对候选 ID 取并集。由于两路原始分数量纲不同，系统先分别进行 Min-Max 归一化，再按 0.45 的向量权重和 0.55 的 BM25 权重计算融合分。BM25 略高的权重来自当前中文评测集的消融结果；配置值在运行时会裁剪为非负数并重新归一化，因此可以通过配置调整，而不需要修改检索代码。"),
        ("Normal", "融合结果随后执行本地可解释重排。最终得分由 0.65 的融合分、0.15 的 Query 覆盖率、0.12 的中文短语命中率和 0.08 的词面相似度组成，不再调用大模型或外部 Reranker。排序第一的 Chunk 会与同文档前后相邻块合并，并识别重叠文本，减少固定切块造成的语义断裂；已被扩展覆盖的相邻块不再作为独立结果重复返回。完成来源去重后默认保留 3 条上下文，并附带来源文件、章节、融合分、重排分和扩展块信息。"),
        ("Normal", "检索结果以“事实参考”而非“可执行指令”注入 ContextBundle。ResponseAgent 可以据此解释常见心理现象、给出非诊断性的支持建议或引用校内资源，但只有证据中明确存在的电话、办公时间、人员和政策才允许出现在回复中。知识文档中的提示性文字不能覆盖系统 Prompt，任何要求忽略规则、输出后台信息或执行工具的内容都不会被当作指令。RAG 与 Skill 的职责保持分离：RAG 回答“有哪些可信事实和资源”，Skill 规定“在当前场景中如何表达以及不能表达什么”。"),
        ("Normal", "检索链路保留 vector、bm25、hybrid、reranked 和 final 五个阶段，便于评测召回、融合、重排和邻块扩展分别带来的变化。业务运行通常只使用 final，管理和测试可以通过阶段结果定位问题。向量库、BM25 索引或知识文件不可用时，系统记录 rag_fallback 并降级到少量内置审核知识；知识不可用不会阻断基础对话或风险安全回复，也不能成为编造资源的理由。知识库更新后应重新构建索引，并通过主题命中率、来源准确性、无结果比例和安全场景覆盖等指标进行回归验证。"),

        ("Heading 2", "5.5 心理场景 Skill"),
        ("Normal", "心理场景 Skill 用于沉淀可复用的表达流程和禁止项，将原本散布在基础 Prompt 中的场景规则组织为 skills/<skill_name>/SKILL.md。每个 Skill 采用“一目录一文件”结构，目录名使用小写字母、数字和下划线；YAML frontmatter 只保存 name 与 description，正文必须包含大小写固定的二级标题 ## Workflow。运行时只读取 Workflow 之后的内容，并将其作为相应场景中的动作步骤、语气要求和禁止项。该形式便于独立审阅、版本管理、自动校验和专项测试，也避免基础 Prompt 因不断叠加场景规则而膨胀。"),
        ("Normal", "Skill 的职责边界是规定“在特定心理场景中怎样回应”，不负责风险定级、任务调度、记忆读取、RAG 检索、数据库访问或工具调用。可变化的校园事实继续由 RAG 提供；心理报告、台账、个案和预警由 AgentHarness 与后台工具链处理；固定安全规则始终具有最高优先级。因此 Skill 不能把日常交流改判为心理支持或风险处置，不能降低 hard_rule 给出的风险下限，也不能自行决定是否创建报告、个案或预警任务。"),
        ("Normal", "系统内置 7 个 Skill：supportive_response_baseline 规定心理支持的基础语气、具体共情和少量行动建议；referral_resource_guidance 将现实求助表达为可执行的下一步并禁止编造资源；anxiety_grounding_support 提供焦虑或惊恐场景下的简单稳定动作；sleep_routine_support 提供当晚可实施的非药物睡眠安排；academic_stress_planning 将学业压力拆成一个任务、一个时间块和一个停止点；high_risk_safety_plan 聚焦未来十至十五分钟的安全、现实陪伴和危险条件隔离；counselor_handoff_summary 以固定字段生成供工作人员查看的事实化交接摘要。"),
        ("Normal", "Skill 路由由 SkillRegistry 根据最终风险等级和 UnderstandingAgent 提取的主题确定。日常交流不加载心理 Skill，也不触发心理 RAG、报告或风险工具；心理支持固定加载 supportive_response_baseline 和 referral_resource_guidance，并按焦虑、睡眠、学业主题叠加相应 Skill，同一轮可以加载多个主题规则，同时生成心理报告、交接摘要和台账，但不创建个案；风险处置加载 referral_resource_guidance 和 high_risk_safety_plan，不再叠加一般焦虑、睡眠或学业建议分散安全重点，同时生成报告、交接摘要、台账和个案，预警任务依赖个案创建成功。"),
        ("Normal", "SkillRegistry 在 Runtime 初始化时扫描 skills 目录，并校验目录命名、SKILL.md 是否存在、UTF-8 编码、YAML frontmatter 是否闭合、name 与目录是否一致、description 是否有效以及 ## Workflow 是否存在。校验结果分为 READY、WARN 和 FAILED：READY 表示结构和内容通过校验；WARN 表示存在非阻断问题但仍可加载；FAILED 表示缺少关键结构、命名错误或文件无法读取，不得进入 Prompt。每个结果同时暴露 issues，管理员可以在“系统运行—Skills”或 GET /api/v1/admin/skills 中查看具体问题。"),
        ("Normal", "Runtime 在进程启动时形成 Skill 快照，保证同一运行周期内使用的规则不会无版本漂移；管理员状态接口每次请求重新扫描磁盘，文件问题可以即时显示，但修改后的 Workflow 必须重启应用才会进入回复链路。新增或修改 Skill 时，需要同步更新必需资产清单、确定性路由、管理端状态和自动化测试，并覆盖 ContextAgent 注入、ResponseAgent 生成、SafetyAgent 审查以及关键场景回退。Skill 的加载状态、路由结果和实际注入列表均写入 Trace，便于审计某轮回复使用了哪些场景规则。"),
        ("Normal", "关键场景设置独立的安全降级门。心理支持缺少 supportive_response_baseline 时，ResponseAgent 不进行无约束自由生成，而是使用固定心理支持回复；风险处置缺少 high_risk_safety_plan 时，直接使用固定高风险安全回复；固定回复仍需经过 SafetyAgent 的确定性审查。counselor_handoff_summary 不进入学生端 Prompt，AgentHarness 仅在心理报告取得 ID 后渲染并写入 psychological_reports.handoff_summary，字段包括报告 ID、学生身份、风险等级、情绪标签、置信度、模型摘要、建议跟进动作和必要长度的学生原始表达摘录。该摘要只供授权工作人员查看，不构成临床诊断，也不得添加学生未表达的事实。"),
    ]

    for style, text in content:
        insert_before(anchor, text, style)

    # Ask Word to refresh the table of contents and other fields on open.
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    expected_headings = [
        "5 模型、上下文与知识设计",
        "5.1 Prompt 与结构化输出",
        "5.2 模型配置与调用",
        "5.3 上下文与分层记忆",
        "5.4 混合 RAG 检索",
        "5.5 心理场景 Skill",
    ]
    reopened = Document(OUTPUT)
    headings = [
        paragraph.text.strip()
        for paragraph in reopened.paragraphs
        if paragraph.text.strip()
        and paragraph.style
        and paragraph.style.name.startswith("Heading")
    ]
    chapter_headings = headings[headings.index(expected_headings[0]):headings.index("6 安全、工具与业务闭环")]
    if chapter_headings != expected_headings:
        raise AssertionError(chapter_headings)
    with ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise AssertionError("Invalid DOCX package")
    print(f"output={OUTPUT}")
    print(f"size={OUTPUT.stat().st_size}")
    print(f"paragraphs={len(reopened.paragraphs)} tables={len(reopened.tables)} images={len(reopened.inline_shapes)}")
    print(f"chapter5_headings={chapter_headings}")


if __name__ == "__main__":
    main()
