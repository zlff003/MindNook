from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document


ROOT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\skills_docs_updated")
SUSPECT = re.compile(r"crisis_response|general_support|风险候选回复回复|5\.6 风险识别|5\.7 管理、审计|4\.1 CollaborationBlackboard设计|4\.6协作")


for path in sorted(ROOT.glob("*.docx")):
    doc = Document(path)
    all_text = [p.text for p in doc.paragraphs]
    all_text += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    joined = "\n".join(all_text)
    headings = [(i, p.style.name, p.text.strip()) for i, p in enumerate(doc.paragraphs) if p.text.strip() and p.style and p.style.name.startswith("Heading")]
    suspicious = [(i, p.text) for i, p in enumerate(doc.paragraphs) if SUSPECT.search(p.text)]
    with ZipFile(path) as z:
        zip_error = z.testzip()
    print("\n", path.name)
    print({"size": path.stat().st_size, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "images": len(doc.inline_shapes), "zip_error": zip_error})
    print({"Skill": joined.count("Skill"), "七个": joined.count("7 个") + joined.count("七个"), "handoff": joined.count("handoff_summary") + joined.count("交接摘要"), "suspicious": suspicious})
    if "设计方案书" in path.name:
        part2 = []
        in_part2 = False
        for i, style, text in headings:
            if text == "1 智能体设计目标":
                in_part2 = True
            if text == "1 测试概述":
                in_part2 = False
            if in_part2:
                part2.append((i, style, text))
        for row in part2:
            print(*row, sep="\t")
    required = [
        "日常交流不触发心理 RAG",
        "心理支持生成心理报告和台账",
        "风险处置生成报告",
        "预警任务依赖个案创建成功",
    ]
    print("rules", {x: (x in joined) for x in required})
