from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\Projects\Agent\MindBridge\项目提交文档\MindNook智能体 - 设计方案书.docx")
STAGE = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\chapter5_rewrite")
OUTPUT = STAGE / SOURCE.name


def find(doc, exact):
    for p in doc.paragraphs:
        if p.text.strip() == exact:
            return p
    raise ValueError(exact)


def insert_before(anchor, text, style="Normal"):
    node = OxmlElement("w:p")
    anchor._p.addprevious(node)
    p = Paragraph(node, anchor._parent)
    p.style = style
    p.add_run(text)
    return p


def delete_range(doc, start_text, end_text):
    start = find(doc, start_text)
    end = find(doc, end_text)
    node = start._p
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def main():
    if STAGE.exists():
        shutil.rmtree(STAGE)
    STAGE.mkdir(parents=True)
    doc = Document(SOURCE)

    # Replace old chapter 5 and standalone chapter 6, stopping before old chapter 7.
    delete_range(doc, "5 模型、上下文与知识设计", "7 安全、工具与业务闭环")
    anchor = find(doc, "7 安全、工具与业务闭环")

    content = [
        ("Heading 1", "5 模型、上下文、知识与 Skill 设计"),
        ("Normal", "本章说明多智能体推理所依赖的模型、Prompt、上下文、分层记忆、知识检索和心理场景 Skill 如何形成统一的回复支撑体系。整体设计遵循“模型负责理解与生成、记忆提供必要连续性、RAG 提供可追溯事实、Skill 提供场景表达规则、安全机制设置不可绕过的边界”的原则。各部分由 ContextAgent 汇总为结构化 ContextBundle，再由 ResponseAgent 形成候选回复，并由 SafetyAgent 在公开前完成复审。"),

        ("Heading 2", "5.1 总体设计与职责边界"),
        ("Normal", "模型、上下文、知识与 Skill 并不是四套彼此独立的能力，而是按照确定的顺序共同服务于一次回复生成。UnderstandingAgent 首先识别对话类型、主题和知识需求；SafetyAgent 根据固定安全规则和当前表达确定风险下限；ContextAgent 再选择需要装配的记忆、知识证据、场景 Skill 和回复约束；ResponseAgent 最后基于这些结构化输入生成候选回复。任何知识或 Skill 都不能直接改变风险等级，也不能绕过 SafetyAgent 的审查结论。"),
        ("Normal", "系统将稳定规则与可变化内容分开维护：Agent 的角色、输入输出协议、隐私边界和危险内容禁令属于基础 Prompt；校园资源、心理知识和事实依据属于 RAG；不同心理场景中的表达步骤和禁止项属于 Skill Workflow；风险定级和工具副作用仍由固定规则、Agent Runtime 与 AgentHarness 控制。该边界可以避免把大量场景规则反复堆入 Prompt，也减少知识事实、表达策略和业务调度相互污染。"),

        ("Heading 2", "5.2 Prompt 分层与结构化输出"),
        ("Normal", "系统 Prompt 采用分层装配方式。第一层是不可覆盖的通用约束，包括角色职责、输入输出边界、不得进行临床诊断、不得提供药物或危险细节、不得泄露后台等级和不得编造校园资源等要求；第二层是本轮业务状态，包括 intent、risk、memoryBrief、modelHistory 和 RAG 证据；第三层是 ContextAgent 按等级和主题选出的 Skill Workflow；第四层是具体任务指令，例如生成候选回复、评估风险或复审指定版本的回复。后加载的知识和 Skill 只能补充内容，不能覆盖上层安全规则。"),
        ("Normal", "UnderstandingAgent 和 SafetyAgent 使用结构化结果，避免自由文本被后续组件误读。IntentAnalysis 保存对话类型、主题、知识需求、置信度和判断理由；风险结果保存分级、情绪概况、即时危险和是否需要安全问题；ResponseReview 保存 approved、violations 和 revision_guidance。ResponseAgent 输出的正文先作为 response_proposal 写入黑板，并明确关联版本号。只有 SafetyAgent 的审查结果与当前 proposal 版本一致时，CoordinatorAgent 才允许采纳。"),
        ("Normal", "Prompt 中只放入完成当前任务所需的信息。学生原始敏感信息在进入模型前完成脱敏；工作人员交接摘要、数据库实体、工具执行状态和其他 Agent 的私有记忆不会进入学生回复 Prompt。模型也不得自行补写未加载 Skill 的规则或知识库没有提供的事实。通过这种最小化装配方式，可以控制 Token 规模，降低提示注入、后台信息泄露和历史判断固化的风险。"),

        ("Heading 2", "5.3 模型配置、调用预算与失败处理"),
        ("Normal", "模型层通过统一的 ModelProvider 接口提供能力，支持 Mock、本地模型和 OpenAI-compatible API，并允许 UnderstandingAgent、SafetyAgent 和 ResponseAgent 按职责使用不同模型或参数。结构化判断任务优先使用低随机性配置，以保证字段稳定；学生回复生成可以保留适度语言自然度，但仍受长度、建议数量和安全约束限制。模型实现可以替换，而 AgentTask、AgentArtifact、ContextBundle 和审查协议保持不变。"),
        ("Normal", "CoordinatorAgent 统一控制单轮协作预算，包括最大协作轮数、模型调用总次数、单次模型超时和整轮截止时间。模型调用前检查剩余预算，调用完成后记录模型、耗时、Token 估算、成功状态和错误摘要。回复修复最多执行一次，防止“生成—拒绝—重写”无限循环；审查结果必须针对当前候选版本，旧版本的通过结论不能复用到新候选。"),
        ("Normal", "模型超时、返回格式错误、结构化解析失败或预算耗尽时，系统不继续开放式尝试。理解任务根据固定规则形成降级 intent；明确危险表达继续保持风险处置下限；回复生成失败时使用带 trusted_fallback 标记的固定模板；安全审查不可用时，心理支持和风险处置进入固定安全回退。所有降级过程写入 AgentStep 和 Trace，保证故障不会被伪装成正常模型结果。"),

        ("Heading 2", "5.4 上下文装配与分层记忆"),
        ("Normal", "AgentHarness 在 Runtime 启动前构造 MemoryContext。MySQL 保存当前会话的完整业务消息，是可回填的事实来源；Redis 使用会话级 Key 缓存最近最多 40 条脱敏消息并设置 TTL。Redis 未命中、内容损坏或服务不可用时，系统从 MySQL 查询当前会话最近的已完成消息，恢复时间顺序后继续运行，并在条件允许时回填缓存。不同学生和不同会话始终使用独立 Key，不允许用其他会话内容补全缺失上下文。"),
        ("Normal", "长对话采用“摘要 + 最近原文窗口”策略。较早历史通过确定性规则压缩为最多 500 字的 memoryBrief，优先保留学生近期关注、系统已经提供的支持和本轮输入关注；最近最多 8 条消息作为 modelHistory 保留原角色和脱敏原文，同时受总字符预算限制。短历史不重复注入摘要，长历史才将 memoryBrief 作为独立 system 背景加入 Prompt，从而兼顾连续性、成本和当前消息优先级。"),
        ("Normal", "各 Agent 的私有记忆使用 agent:<AgentName>:<session_public_id> 隔离保存。UnderstandingAgent 的历史意图、SafetyAgent 的风险判断、ContextAgent 的检索策略和 ResponseAgent 的回复策略只供本角色参考，不能跨角色直接注入。私有记忆不能覆盖当前输入和固定规则，后台风险标签也不能作为学生的长期固定画像。完整 Agent Run、步骤、黑板事件和检索信息写入 Trace 供审计，但 Trace 永不回流下一轮 Prompt。"),

        ("Heading 2", "5.5 RAG 触发、检索与证据注入"),
        ("Normal", "RAG 根据最终分级按需触发。日常交流只装配必要的近期上下文，不调用心理 RAG，避免把学习、编程或校园事务强行心理化；心理支持检索与当前困扰相关的审核知识；风险处置在正常 Query 和主题之外，固定补充“当前安全、危机支持、现实求助”等检索主题。RAG 只提供知识证据，不参与最终风险定级，也不能降低固定规则确定的安全下限。"),
        ("Normal", "知识文件采用 UTF-8 Markdown，按章节和字符边界切分，默认块大小为 512 字符、重叠 64 字符。在线检索同时执行 Chroma 向量召回和 BM25 关键词召回，两路候选经分数归一化与加权融合后，再使用本地可解释特征重排。第一名结果会合并同文档相邻块，减少固定切分造成的语义断裂；最终按来源去重后向 ContextAgent 返回最多 3 条带文件名和章节信息的证据。"),
        ("Normal", "ContextAgent 将检索结果作为“事实参考”注入 ContextBundle，并同时加入不得编造资源、不得输出后台标签和风险场景优先现实安全行动等硬约束。知识文档中的文本不能被当作系统指令执行，校园电话、人员、办公时间和政策只有在检索上下文明示时才可使用。检索失败时系统降级到少量内置审核知识并记录 rag_fallback；知识不可用不得阻断基础对话，也不得成为编造资源的理由。"),

        ("Heading 2", "5.6 心理场景 Skill 资产设计"),
        ("Normal", "心理场景 Skill 将原本分散在基础 Prompt 中的表达规则沉淀为 skills/<skill_name>/SKILL.md。每个 Skill 采用“一目录一文件”结构，目录使用小写下划线命名；YAML frontmatter 只保存 name 和 description，正文必须包含固定标题“## Workflow”。运行时只读取 Workflow 之后的内容，将其作为本场景需要遵循的动作步骤和禁止项。该结构便于独立审阅、版本管理、自动校验和专项测试。"),
        ("Normal", "Skill 只规定“在特定心理场景中应该怎样表达”和“哪些内容不能出现”，不负责任务调度、风险定级、记忆读取、RAG 检索、数据库访问或工具调用。可变化的校园事实继续由 RAG 提供；报告、台账、个案和预警由 AgentHarness 与工具队列处理；固定安全规则仍具有最高优先级。Skill 不能改变业务分级，不能降低风险下限，也不能自行决定是否创建个案或发送预警。"),
        ("Normal", "系统内置 7 个 Skill：supportive_response_baseline 规定心理支持的基础语气和少量行动建议；referral_resource_guidance 负责现实求助与不编造资源；anxiety_grounding_support 提供焦虑稳定动作；sleep_routine_support 提供非药物的睡眠安排；academic_stress_planning 将学业压力拆成下一小步；high_risk_safety_plan 聚焦未来 10 至 15 分钟的安全；counselor_handoff_summary 生成工作人员可见的事实化交接摘要。"),

        ("Heading 2", "5.7 Skill 路由、加载与安全降级"),
        ("Normal", "SkillRegistry 根据最终分级和 UnderstandingAgent 提取的主题执行确定性路由。日常交流不加载心理 Skill；心理支持固定加载 supportive_response_baseline 和 referral_resource_guidance，并根据焦虑、睡眠和学业标记叠加对应主题 Skill，同一轮可以加载多个主题规则；风险处置加载 referral_resource_guidance 和 high_risk_safety_plan，不继续叠加一般焦虑、睡眠或学业建议，以免分散当前安全重点。"),
        ("Normal", "SkillRegistry 启动时扫描配置的 skills 目录，校验目录命名、UTF-8、YAML frontmatter、name 与目录一致性、description 和 Workflow。校验结果分为 READY、WARN 和 FAILED：READY 与 WARN 可以进入 Runtime，FAILED 不注入 Prompt，并通过 issues 暴露具体问题。Runtime 在进程生命周期内使用启动时快照，保证同一次运行中的规则不会无版本漂移；管理员状态接口每次请求重新扫描磁盘，修改文件后仍需重启应用才会生效。"),
        ("Normal", "关键场景设置独立安全门。心理支持缺少 supportive_response_baseline 时，ResponseAgent 不进行无约束自由生成，而是使用固定心理支持回复；风险处置缺少 high_risk_safety_plan 时，系统使用固定高风险安全回复。固定回复仍要经过 SafetyAgent 的确定性检查。Skill 的损坏、缺失或警告会进入管理端状态和 Trace，但不得导致学生直接收到未经约束的敏感回复。"),

        ("Heading 2", "5.8 工作人员交接摘要与 Skill 可观测性"),
        ("Normal", "counselor_handoff_summary 是面向工作人员的报告 Skill，不进入学生端回复 Prompt。AgentHarness 在心理支持或风险处置报告取得 ID 后，按固定字段生成 handoff_summary，包括报告 ID、学生身份、分级结果、情绪标签、置信度、模型摘要、建议跟进动作和必要长度的学生原始表达摘录。心理支持建议工作人员复核困扰、日常影响与支持需求；风险处置优先提示核实当前安全、陪伴情况和现实支持连接。"),
        ("Normal", "风险个案复用报告中的同一份交接摘要，避免报告和个案分别生成造成信息漂移。历史报告缺少摘要时，应用启动过程执行幂等回填，并同步已有风险个案。报告列表默认不直接返回 handoff_summary 和原始表达；只有授权工作人员打开报告上下文时才读取，并记录访问审计。交接摘要只用于人工辅助，不构成临床诊断，也不得添加学生没有表达的事实。"),
        ("Normal", "管理端“系统运行—Skills”和 GET /api/v1/admin/skills 展示每个 Skill 的状态、描述和 issues，便于发现文件缺失、命名错误或 Workflow 不完整等问题。新增或修改 Skill 时，需要同步更新必需资产清单、确定性路由和自动化测试，并验证 ContextAgent 注入、ResponseAgent 生成、SafetyAgent 审查、关键 Skill 回退及报告摘要持久化。通过文件状态、运行 Trace 和专项测试三层证据，确保心理表达规则可维护、可观察、可追溯。"),
    ]

    for style, text in content:
        insert_before(anchor, text, style)

    # Renumber the following compact chapters after merging old chapter 6.
    replacements = {
        "7 安全、工具与业务闭环": "6 安全、工具与业务闭环",
        "7.1 分级与回复安全": "6.1 分级与回复安全",
        "7.2 PII、权限与服务边界": "6.2 PII、权限与服务边界",
        "7.3 工具 DAG 与异步执行": "6.3 工具 DAG 与异步执行",
        "8 接口、数据与可观测性": "7 接口、数据与可观测性",
        "8.1 Runtime 与 Harness 接口": "7.1 Runtime 与 Harness 接口",
        "8.2 核心数据与事件": "7.2 核心数据与事件",
        "8.3 管理端观测": "7.3 管理端观测",
        "9 异常、测试与部署": "8 异常、测试与部署",
        "9.1 异常与降级": "8.1 异常与降级",
        "9.2 测试与验收": "8.2 测试与验收",
        "9.3 部署与当前限制": "8.3 部署与当前限制",
    }
    for old, new in replacements.items():
        set_text(find(doc, old), new)

    doc.save(OUTPUT)
    print(OUTPUT, OUTPUT.stat().st_size, len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes))


if __name__ == "__main__":
    main()
