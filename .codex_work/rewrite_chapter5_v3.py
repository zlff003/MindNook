from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\Projects\Agent\MindBridge\项目提交文档\MindNook智能体 - 设计方案书 - 第5章重构版.docx")
STAGED = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\MindNook智能体 - 设计方案书 - 第5章结构优化版.docx")


def find_paragraph(document, exact_text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact_text}")


def delete_range(document, start_text, end_text):
    start = find_paragraph(document, start_text)
    end = find_paragraph(document, end_text)
    node = start._p
    while node is not None and node is not end._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node


def insert_paragraph(anchor, text="", style="Normal", bold_lead=False):
    node = OxmlElement("w:p")
    anchor._p.addprevious(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = style
    if bold_lead and "：" in text:
        lead, rest = text.split("：", 1)
        run = paragraph.add_run(lead + "：")
        run.bold = True
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)
    return paragraph


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + margin
        element = tc_mar.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    # Match the start cell margin so the visible outer border aligns with body text.
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for column_index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(widths[column_index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def format_table(table):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(10.5)
                    if row_index == 0:
                        run.bold = True


def insert_table(document, anchor, headers, rows, widths, style):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = style
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    set_table_geometry(table, widths)
    format_table(table)
    anchor._p.addprevious(table._tbl)
    return table


def add_heading(anchor, text, level):
    return insert_paragraph(anchor, text, f"Heading {level}")


def add_bullets(anchor, items):
    for item in items:
        insert_paragraph(anchor, item, "List Bullet", bold_lead=True)


def add_steps(anchor, items):
    for item in items:
        insert_paragraph(anchor, item, "List Number", bold_lead=True)


def main():
    document = Document(SOURCE)
    original_table_count = len(document.tables)
    table_style = document.tables[0].style if document.tables else "Table Grid"

    delete_range(document, "5 模型、上下文与知识设计", "6 安全、工具与业务闭环")
    anchor = find_paragraph(document, "6 安全、工具与业务闭环")

    add_heading(anchor, "5 模型、上下文与知识设计", 1)
    insert_paragraph(
        anchor,
        "本章说明 MindNook 如何将模型推理、Prompt 约束、上下文管理、分层记忆、知识检索与心理场景 Skill 组织为统一的回复支撑体系。各部分按照“模型负责理解与生成、上下文负责本轮信息装配、记忆提供受控连续性、RAG 提供可追溯事实、Skill 提供场景表达规则”的边界协同工作，最终由 SafetyAgent 审查候选回复、CoordinatorAgent 决定是否采纳。",
    )
    add_bullets(anchor, [
        "安全边界优先：模型、知识和 Skill 均不能覆盖固定安全规则，也不能自行降低风险等级。",
        "最小必要装配：只向当前 Agent 提供完成任务所需的脱敏信息，避免上下文无边界增长。",
        "结构化协作：意图、风险、上下文、候选回复和审查结果均以版本化 Artifact 交换。",
        "可降级与可追溯：模型、Redis、RAG 或 Skill 异常时进入确定性回退，并将原因写入 Trace。",
    ])

    add_heading(anchor, "5.1 Prompt 与结构化输出", 2)
    insert_paragraph(
        anchor,
        "Prompt 采用分层装配，不把角色设定、历史消息、知识全文和场景规则一次性堆入同一输入。上层规则约束下层内容，后加载的 RAG 证据和 Skill Workflow 只能补充回复依据，不能改写系统边界。",
    )
    insert_table(
        document,
        anchor,
        ["层级", "主要内容", "来源", "控制要求"],
        [
            ["第一层：基础约束", "Agent 职责、隐私边界、危险内容与资源真实性要求", "固定系统 Prompt", "不可被后续内容覆盖"],
            ["第二层：运行状态", "当前输入、intent、risk、memoryBrief、modelHistory", "黑板与 MemoryContext", "仅装配当前任务所需字段"],
            ["第三层：场景依据", "RAG 证据、适用 Skill Workflow、回复约束", "ContextAgent", "知识只作事实参考，Skill 只作表达规则"],
            ["第四层：任务指令", "理解、评估、生成、审查或修复指定版本", "AgentTask", "输出必须符合指定结构"],
        ],
        [1150, 3150, 2050, 2650],
        table_style,
    )
    insert_paragraph(anchor, "各智能体的 Prompt 与输出职责如下：")
    add_bullets(anchor, [
        "UnderstandingAgent：识别对话类型、主题、知识需求、置信度和判断理由；不生成安慰、风险处置或学生回复。",
        "SafetyAgent：风险评估阶段输出风险等级、即时危险和安全问题；回复审查阶段输出 approved、violations 与 revision_guidance。",
        "ContextAgent：确定性装配记忆、RAG 证据、Skill 和回复约束，形成 ContextBundle，不直接生成面向学生的结论。",
        "ResponseAgent：依据 hard_rule、intent、risk 和 ContextBundle 生成带版本号的 response_proposal。",
    ])
    insert_paragraph(anchor, "候选回复按以下顺序完成结构化审查与采纳：")
    add_steps(anchor, [
        "生成候选：ResponseAgent 将完整正文发布为内部 response_proposal，不直接发送给学生。",
        "执行审查：SafetyAgent 同时进行模型审查和确定性规则检查，并明确关联候选版本。",
        "受限修复：首次不通过时依据 violations 和 revision_guidance 生成一次新版本，再次提交审查。",
        "最终采纳：只有当前版本 approved=true 时，CoordinatorAgent 才发布 final_response；再次失败或预算耗尽时使用固定安全回退。",
    ])
    add_bullets(anchor, [
        "隐私最小化：手机号、邮箱、身份证号和学号等信息在进入模型前脱敏。",
        "信息隔离：交接摘要、数据库实体、工具状态和其他 Agent 私有记忆不进入学生端 Prompt。",
        "防提示注入：RAG 文本不得作为系统指令执行，知识库未提供的人员、电话、时间和政策不得推断生成。",
    ])

    add_heading(anchor, "5.2 模型配置与调用", 2)
    insert_paragraph(
        anchor,
        "模型能力通过统一 ModelProvider 接口接入。测试环境使用确定性的 MockModelProvider，部署环境可配置 OpenAI-compatible API；模型标识、服务地址、鉴权、超时与采样参数均由配置注入，业务代码不直接依赖特定厂商 SDK。",
    )
    insert_table(
        document,
        anchor,
        ["任务", "执行角色", "配置侧重点", "失败处理"],
        [
            ["意图理解", "UnderstandingAgent", "低随机性、结构字段稳定", "固定规则生成降级 intent"],
            ["风险评估", "SafetyAgent", "保守判断、保留安全下限", "hard_rule 风险下限继续生效"],
            ["回复生成", "ResponseAgent", "适度自然度、限制篇幅与建议数量", "使用 trusted_fallback 固定模板"],
            ["回复审查", "SafetyAgent", "低随机性、版本绑定、违规项明确", "心理支持与风险场景固定安全回退"],
        ],
        [1500, 1900, 3100, 2500],
        table_style,
    )
    insert_paragraph(anchor, "CoordinatorAgent 对单轮调用设置统一预算：")
    add_bullets(anchor, [
        "最大协作轮数：默认 16 轮。",
        "模型调用上限：单轮最多 7 次。",
        "单次模型超时：默认 20 秒。",
        "整轮模型运行时限：默认 55 秒。",
        "回复修复次数：最多 1 次，防止“生成—拒绝—重写”无限循环。",
    ])
    insert_paragraph(anchor, "每次模型调用前检查剩余预算，调用后记录 Agent、模型、耗时、Token 估算、成功状态和错误摘要。达到任一上限后不再开放式尝试，而是使用已有 Artifact 或固定模板完成本轮。ResponseAgent 产生的流式 Token 先在服务端累积，只有 SafetyAgent 审查通过且 CoordinatorAgent 采纳后，AgentHarness 才通过公开 SSE 按片段发送最终成品。")
    add_bullets(anchor, [
        "结构解析失败：不把自由文本直接交给后续 Agent，改用对应的确定性降级结果。",
        "安全审查不可用：日常交流仅在通过固定输出规则后放行；心理支持与风险处置使用固定安全回复。",
        "调用预算耗尽：停止继续生成，记录预算事件并用现有产物结束本轮。",
        "全过程留痕：超时、解析失败、回退原因和最终采用版本均写入 AgentStep、Event 与 RuntimeTrace。",
    ])

    add_heading(anchor, "5.3 上下文管理与压缩", 2)
    insert_paragraph(
        anchor,
        "上下文管理解决的是“本轮模型应该看到什么”。AgentHarness 先从当前会话的可信消息中构造 MemoryContext，ContextAgent 再结合 intent、risk、RAG 与 Skill 形成 ContextBundle。上下文不是数据库记录的完整复制，而是经过会话隔离、脱敏、排序、截断和压缩后的任务输入。",
    )
    insert_paragraph(anchor, "单轮上下文按照以下流程装配：")
    add_steps(anchor, [
        "确认会话边界：校验学生与会话归属，只读取当前 session_public_id 对应的数据。",
        "加载近期消息：优先读取 Redis 中最近最多 40 条脱敏消息；未命中或不可用时从 MySQL 回填 completed 消息。",
        "恢复与清洗：按时间顺序恢复角色消息，移除未完成记录，并对手机号、邮箱、身份证号和学号等信息脱敏。",
        "压缩长历史：保留最近原文窗口，将较早历史压缩为结构固定的 memoryBrief。",
        "装配任务上下文：根据 Agent 职责加入当前输入、必要 Artifact、RAG 证据、Skill Workflow 和回复约束。",
    ])
    insert_table(
        document,
        anchor,
        ["上下文组件", "默认范围", "内容形态", "主要用途"],
        [
            ["当前输入", "本轮 1 条", "单独保存的脱敏原文", "保持最高关注优先级"],
            ["modelHistory", "最近最多 8 条；总计不超过 6000 字符", "带 user/assistant 角色的脱敏原文", "保留近期措辞和问答连续性"],
            ["memoryBrief", "最多 500 字符", "4 条学生关注 + 3 条助手支持 + 本轮关注", "压缩较早历史并保留主题脉络"],
            ["ContextBundle", "按本轮任务动态选择", "记忆、证据、Skill 与回复约束的结构化集合", "作为 ResponseAgent 的唯一上下文入口"],
        ],
        [1550, 2200, 3100, 2150],
        table_style,
    )
    insert_paragraph(anchor, "长历史压缩遵循以下规则：")
    add_bullets(anchor, [
        "结构优先：memoryBrief 固定保留“学生关注、已提供支持、本轮关注”三个栏目，整体超长时不因直接截断丢失栏目。",
        "近期优先：从最近消息向前选择 modelHistory；单条过长时保留开头和结尾，中间使用省略。",
        "不重复注入：短历史只使用角色化原文；超过原文窗口时才额外加入 memoryBrief。",
        "不制造事实：摘要只复述学生已经表达的关注和系统已经提供的支持，不生成诊断或推断未表达经历。",
        "当前输入独立：本轮消息不依赖摘要表达，避免压缩结果削弱新信息的优先级。",
    ])
    insert_paragraph(anchor, "完整 Agent Run、协作步骤、检索阶段和模型调用信息只进入 RuntimeTrace，用于审计与问题定位，不回流下一轮 Prompt。最终事务提交成功后，学生消息与助手成品才追加到 Redis；事务回滚时不写入未完成回复，保证上下文与学生实际收到的内容一致。")

    add_heading(anchor, "5.4 分层记忆设计", 2)
    insert_paragraph(
        anchor,
        "分层记忆解决的是“信息保存在哪里、保留多久、由谁使用”。MindNook 将完整业务记录、短期会话窗口、压缩摘要、Agent 私有记忆和运行追踪分开管理，使会话连续性、性能、隐私隔离和可审计性之间保持平衡。",
    )
    insert_table(
        document,
        anchor,
        ["记忆层级", "存储或载体", "保留策略", "使用边界"],
        [
            ["业务消息层", "MySQL chat_messages", "保存完整业务消息和完成状态", "会话记录的事实来源；不等同于全部进入模型"],
            ["短期会话层", "Redis short-term-memory", "最近最多 40 条；TTL 86400 秒", "仅限当前学生、当前会话；失效后可由 MySQL 回填"],
            ["压缩摘要层", "MemoryContext.memoryBrief", "本轮动态生成；最多 500 字符", "只服务当前运行，不作为长期用户画像"],
            ["角色私有层", "Redis agent:<AgentName>:<session_id>", "每个 Key 最多 20 条；TTL 86400 秒；Prompt 读取最近 4 条", "仅本 Agent 可读，禁止跨角色直接注入"],
            ["运行追踪层", "RuntimeTrace / AgentStep / Event", "随运行与审计记录保存", "用于排障和审计，永不进入下一轮 Prompt"],
        ],
        [1450, 2500, 2500, 2550],
        table_style,
    )
    insert_paragraph(anchor, "各层之间必须遵守以下隔离规则：")
    add_bullets(anchor, [
        "会话隔离：不同学生、不同会话使用独立 Key，不得使用其他会话补全当前上下文。",
        "角色隔离：UnderstandingAgent 的历史意图、SafetyAgent 的风险判断、ContextAgent 的检索策略和 ResponseAgent 的表达策略只供本角色参考。",
        "风险去标签化：后台风险等级和历史风险判断不能固化为学生的长期画像，也不能覆盖当前输入。",
        "事实来源明确：MySQL 是完整业务消息的可信来源；Redis 是性能缓存，MemoryContext 是本轮派生视图。",
        "失败不越界：Redis 不可用时只从当前会话的 MySQL 消息回填，不扩大数据范围，也不跳过脱敏。",
    ])
    insert_paragraph(anchor, "记忆写入与恢复采用受控顺序：")
    add_steps(anchor, [
        "读取：优先 Redis，异常时从 MySQL 查询当前会话最近 40 条 completed 消息。",
        "恢复：按时间重新排序、脱敏并回填 Redis，刷新 TTL。",
        "使用：分别生成 modelHistory、memoryBrief 和角色私有上下文，按 Agent 职责最小化注入。",
        "写回：数据库事务成功后追加本轮学生消息和最终回复，并将 Redis 裁剪至 40 条。",
    ])

    add_heading(anchor, "5.5 混合 RAG 检索", 2)
    insert_paragraph(
        anchor,
        "心理知识检索由 ContextAgent 按最终风险等级触发。日常交流不调用心理 RAG；心理支持根据当前困扰、主题和知识需求构造 Query；风险处置除当前输入外，固定补充“当前安全、危机支持、现实求助”等安全主题。RAG 只提供知识证据，不参与最终风险定级，也不能降低固定规则确定的安全下限。",
    )
    insert_paragraph(anchor, "混合检索链路如下：")
    add_steps(anchor, [
        "知识切分：UTF-8 Markdown 按章节和自然语言边界切块，保留文件名、章节路径和块序号。",
        "双路召回：Chroma 负责语义向量召回，BM25 负责关键词与专有名词召回，两路各取最多 12 个候选。",
        "分数归一：对两路原始分分别执行 Min-Max 归一化，消除量纲差异。",
        "加权融合：按向量 0.45、BM25 0.55 计算融合分，权重可配置并自动归一。",
        "本地重排：综合融合分、Query 覆盖率、中文短语命中率和词面相似度重新排序。",
        "邻块扩展：第一名结果合并同文档前后块，去除 overlap 重复，并按来源去重后返回最多 3 条证据。",
    ])
    insert_table(
        document,
        anchor,
        ["参数", "默认值", "设计说明"],
        [
            ["Chunk 大小", "512 字符", "控制单块语义完整性与检索粒度"],
            ["Chunk 重叠", "64 字符", "减少切分边界的信息损失"],
            ["candidate_k", "每路 12 条", "先扩大候选范围，再进行融合和重排"],
            ["融合权重", "向量 0.45 / BM25 0.55", "兼顾语义相似与中文关键词稳定命中"],
            ["重排权重", "0.65 / 0.15 / 0.12 / 0.08", "依次对应融合分、覆盖率、短语命中和词面相似"],
            ["最终结果", "最多 3 条", "完成邻块扩展与来源去重后注入"],
        ],
        [2100, 2200, 4700],
        table_style,
    )
    add_bullets(anchor, [
        "证据注入：检索结果以“事实参考”进入 ContextBundle，并携带来源文件、章节、融合分、重排分和扩展块信息。",
        "资源真实性：电话、办公时间、人员和政策只有在证据明确存在时才允许输出。",
        "职责分离：RAG 回答“有哪些可信事实与资源”，Skill 规定“当前场景如何表达以及不能表达什么”。",
        "可观测性：保留 vector、bm25、hybrid、reranked 和 final 五个阶段，业务使用 final，评测可比较各阶段变化。",
        "失败降级：索引或知识文件不可用时记录 rag_fallback，并使用少量内置审核知识；不得因此阻断风险安全回复或编造资源。",
    ])

    add_heading(anchor, "5.6 心理场景 Skill", 2)
    insert_paragraph(
        anchor,
        "心理场景 Skill 用于沉淀可复用的表达流程和禁止项。每个 Skill 使用 skills/<skill_name>/SKILL.md，目录采用小写下划线命名；YAML frontmatter 只包含 name 和 description，正文必须包含 ## Workflow。运行时仅将 Workflow 之后的内容作为动作步骤、语气要求和禁止项注入，不把 Skill 当作风险判定、知识库或业务工具。",
    )
    insert_table(
        document,
        anchor,
        ["Skill", "适用场景", "核心作用"],
        [
            ["supportive_response_baseline", "所有心理支持回复", "规定具体共情、温和语气和一至两个小步骤"],
            ["referral_resource_guidance", "心理支持与风险处置", "把现实求助转成可执行下一步，禁止编造资源"],
            ["anxiety_grounding_support", "焦虑、惊恐、呼吸紧张", "先承认难受，再提供一个简单稳定动作"],
            ["sleep_routine_support", "失眠、熬夜、作息紊乱", "提供当晚可实施的一至两个非药物安排"],
            ["academic_stress_planning", "考试、绩点、论文和作业压力", "拆成一个任务、一个时间块和一个停止点"],
            ["high_risk_safety_plan", "所有风险处置场景", "聚焦未来十至十五分钟安全、现实陪伴和危险条件隔离"],
            ["counselor_handoff_summary", "心理支持与风险报告", "生成供工作人员查看的事实化交接摘要"],
        ],
        [3100, 2450, 3450],
        table_style,
    )
    insert_paragraph(anchor, "Skill 路由与业务等级保持一致：")
    add_bullets(anchor, [
        "日常交流：不加载心理 Skill，不触发心理 RAG、报告或风险工具。",
        "心理支持：固定加载 supportive_response_baseline 和 referral_resource_guidance，按焦虑、睡眠、学业主题叠加相应 Skill；生成心理报告、交接摘要和台账，不创建个案。",
        "风险处置：加载 referral_resource_guidance 和 high_risk_safety_plan，不叠加一般主题建议；生成报告、交接摘要、台账和个案，预警任务依赖个案创建成功。",
    ])
    insert_paragraph(anchor, "SkillRegistry 的扫描、校验与加载流程如下：")
    add_steps(anchor, [
        "扫描资产：Runtime 初始化时扫描 skills 目录，并将 7 个内置 Skill 作为必需资产。",
        "执行校验：检查目录命名、SKILL.md、UTF-8、YAML frontmatter、name 与目录一致性、description 和 ## Workflow。",
        "标记状态：READY 和 WARN 可以注入，FAILED 不进入 Prompt；每项问题通过 issues 暴露。",
        "形成快照：Runtime 在当前进程中使用启动快照，避免规则无版本漂移；修改 Workflow 后需重启应用。",
        "提供观测：管理员可在“系统运行—Skills”或 GET /api/v1/admin/skills 查看实时扫描状态。",
    ])
    add_bullets(anchor, [
        "心理支持关键回退：缺少 supportive_response_baseline 时不进行无约束自由生成，改用固定心理支持回复。",
        "风险处置关键回退：缺少 high_risk_safety_plan 时使用固定高风险安全回复；固定回复仍需确定性审查。",
        "交接摘要隔离：counselor_handoff_summary 不进入学生端 Prompt，只在报告取得 ID 后写入 handoff_summary，且不得添加学生未表达的事实。",
        "能力边界：Skill 不能修改风险等级、执行 RAG、访问数据库或决定报告、台账、个案和预警是否创建。",
    ])

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    STAGED.parent.mkdir(parents=True, exist_ok=True)
    document.save(STAGED)

    reopened = Document(STAGED)
    expected_headings = [
        "5 模型、上下文与知识设计",
        "5.1 Prompt 与结构化输出",
        "5.2 模型配置与调用",
        "5.3 上下文管理与压缩",
        "5.4 分层记忆设计",
        "5.5 混合 RAG 检索",
        "5.6 心理场景 Skill",
    ]
    heading_texts = [
        paragraph.text.strip()
        for paragraph in reopened.paragraphs
        if paragraph.text.strip() and paragraph.style and paragraph.style.name.startswith("Heading")
    ]
    chapter_headings = heading_texts[
        heading_texts.index(expected_headings[0]):heading_texts.index("6 安全、工具与业务闭环")
    ]
    if chapter_headings != expected_headings:
        raise AssertionError(chapter_headings)
    if len(reopened.tables) != original_table_count + 6:
        raise AssertionError((original_table_count, len(reopened.tables)))
    with ZipFile(STAGED) as archive:
        if archive.testzip() is not None:
            raise AssertionError("Invalid DOCX package")
    print(f"output={STAGED}")
    print(f"size={STAGED.stat().st_size}")
    print(f"paragraphs={len(reopened.paragraphs)} tables={len(reopened.tables)} images={len(reopened.inline_shapes)}")
    print(f"chapter5_headings={chapter_headings}")


if __name__ == "__main__":
    main()
