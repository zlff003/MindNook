from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document


SOURCE_ROOT = Path(r"D:\Projects\Agent\MindBridge\项目提交文档")
STAGE_ROOT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\localized_route_terms")


REPLACEMENTS = (
    ("CHAT / CONSULT / RISK", "日常交流 / 心理支持 / 风险处置"),
    ("CHAT/CONSULT/RISK", "日常交流/心理支持/风险处置"),
    ("CHAT、CONSULT、RISK", "日常交流、心理支持、风险处置"),
    ("CHAT、CONSULT 或 RISK", "日常交流、心理支持或风险处置"),
    ("CHAT、CONSULT或RISK", "日常交流、心理支持或风险处置"),
    ("CHAT、CONSULT", "日常交流、心理支持"),
    ("CONSULT 与 RISK", "心理支持与风险处置"),
    ("CONSULT和RISK", "心理支持和风险处置"),
    ("CONSULT 或 RISK", "心理支持或风险处置"),
    ("CONSULT或RISK", "心理支持或风险处置"),
    ("RISK 预警", "风险预警"),
    ("RISK 安全响应", "风险安全响应"),
    ("RISK 安全检查", "风险回复安全检查"),
    ("RISK 安全", "风险安全"),
    ("RISK 场景", "风险场景"),
    ("RISK 事件", "风险事件"),
    ("RISK 任务", "风险任务"),
    ("RISK 回复", "风险回复"),
    ("RISK 候选", "风险候选回复"),
    ("RISK 识别", "风险识别"),
    ("RISK 覆盖", "风险覆盖"),
    ("RISK 处理", "风险处理"),
    ("RISK 跟进", "风险跟进"),
    ("RISK 判定", "风险处置判定"),
    ("RISK 个案", "风险个案"),
    ("RISK 报告", "风险处置报告"),
    ("RISK 路径", "风险处置路径"),
    ("CONSULT 场景", "心理支持场景"),
    ("CONSULT 回答", "心理支持回答"),
    ("CONSULT 支持性回答", "心理支持性回答"),
    ("CHAT 普通回答", "日常交流回答"),
    ("CHAT 候选", "日常交流候选回复"),
    ("CHAT 不得", "日常交流不得"),
    ("CHAT 不", "日常交流不"),
    ("CONSULT 不", "心理支持不"),
    ("RISK 不", "风险处置不"),
    ("CONSULT", "心理支持"),
    ("CHAT", "日常交流"),
    ("RISK", "风险处置"),
)


def replace_runs(paragraph) -> int:
    changed = 0
    for run in paragraph.runs:
        before = run.text
        after = before
        for old, new in REPLACEMENTS:
            after = after.replace(old, new)
        if after != before:
            run.text = after
            changed += 1
    return changed


def paragraphs(document: Document):
    seen = set()
    for paragraph in document.paragraphs:
        if paragraph._p not in seen:
            seen.add(paragraph._p)
            yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph._p not in seen:
                        seen.add(paragraph._p)
                        yield paragraph
    for section in document.sections:
        for region in (section.header, section.footer):
            for paragraph in region.paragraphs:
                if paragraph._p not in seen:
                    seen.add(paragraph._p)
                    yield paragraph
            for table in region.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            if paragraph._p not in seen:
                                seen.add(paragraph._p)
                                yield paragraph


def post_fixes(document: Document) -> int:
    fixes = (
        ("风险处置风险", "风险处置"),
        ("心理支持支持", "心理支持"),
        ("日常交流交流", "日常交流"),
        ("风险处置安全响应", "风险安全响应"),
        ("风险处置安全检查", "风险回复安全检查"),
        ("风险处置安全回退", "风险安全回退"),
        ("风险处置安全优先", "风险安全优先"),
        ("风险处置处置", "风险处置"),
        ("风险处置个案", "风险个案"),
        ("风险处置预警", "风险预警"),
        ("心理支持报告", "心理报告"),
        ("日常交流_support", "general_support"),
        ("风险处置_level", "risk_level"),
        ("风险处置 流程", "风险处置流程"),
        ("心理支持与风险处置 得到", "心理支持与风险处置得到"),
        ("提供可执行的 心理支持 支持", "提供可执行的心理支持"),
        ("优先处理 风险处置", "优先处理风险处置"),
        ("风险处置 生成", "风险处置生成"),
        ("风险处置 信息", "风险处置信息"),
        ("风险处置 闭环", "风险处置闭环"),
        ("风险处置 场景", "风险场景"),
        ("风险处置场景", "风险场景"),
        ("心理支持 场景", "心理支持场景"),
        ("心理支持 生成", "心理支持生成"),
        ("心理支持 保存", "心理支持保存"),
        ("风险处置 保存", "风险处置保存"),
        ("心理支持 回复", "心理支持回复"),
        ("心理支持 的", "心理支持的"),
        ("风险处置 的", "风险处置的"),
        ("心理支持 只", "心理支持只"),
        ("为 心理支持与风险处置", "为心理支持与风险处置"),
        ("日常交流 无", "日常交流无"),
        ("风险处置 与", "风险处置与"),
        ("风险处置 对话", "风险处置对话"),
        ("风险处置 候选", "风险处置候选"),
        ("风险处置 优先", "风险处置优先"),
        ("必须为 风险处置", "必须为风险处置"),
        ("降级为 日常交流 或 心理支持", "降级为日常交流或心理支持"),
        ("按 日常交流、心理支持、风险处置 等级", "按分级结果"),
        ("复核 日常交流、心理支持、风险处置 等级", "复核分级结果"),
        ("根据 日常交流、心理支持、风险处置 等级", "根据日常交流、心理支持、风险处置三个等级"),
        ("日常交流、心理支持、风险处置 等级", "日常交流、心理支持、风险处置等级"),
        ("日常交流/心理支持/风险处置 等级", "日常交流/心理支持/风险处置等级"),
        ("最终 日常交流/心理支持/风险处置等级", "最终分级结果"),
        ("日常交流、心理支持、风险处置 三个等级", "日常交流、心理支持、风险处置三个等级"),
        ("日常交流/心理支持/风险处置 三等级", "日常交流/心理支持/风险处置三个等级"),
        ("按 日常交流、心理支持、风险处置 驱动", "按日常交流、心理支持、风险处置驱动"),
        ("真实风险处置若被漏判", "真实风险信号若被漏判"),
        ("明确风险处置召回率", "明确风险信号召回率"),
        ("风险处置召回率", "风险场景召回率"),
        ("风险处置错误降级率", "风险场景错误降级率"),
        ("风险处置夜间值守", "风险场景夜间值守"),
        ("风险处置表达", "风险表达"),
        ("风险处置固定规则", "风险固定规则"),
        ("风险处置优先当前安全", "风险场景优先当前安全"),
        ("固定风险处置规则", "固定安全规则"),
        ("系统按 日常交流", "系统按日常交流"),
        ("提供 日常交流回答", "提供日常交流回答"),
        ("或 风险安全优先回复", "或风险安全优先回复"),
        ("心理支持或风险处置 处理路径", "心理支持或风险处置处理路径"),
        ("心理支持与风险处置 消息", "心理支持与风险处置消息"),
        ("仅 风险处置报告", "仅风险处置报告"),
        ("包括 日常交流", "包括日常交流"),
        ("覆盖 日常交流", "覆盖日常交流"),
        ("学校 风险处置判定", "学校风险处置判定"),
        ("进入 日常交流", "进入日常交流"),
        ("心理支持 调用", "心理支持调用"),
        ("风险处置 模式", "风险处置模式"),
        ("心理支持或风险处置 处理", "心理支持或风险处置处理"),
        ("心理支持与风险处置 执行", "心理支持与风险处置执行"),
        ("采用 日常交流", "采用日常交流"),
        ("心理支持 自动", "心理支持自动"),
        ("日常交流、心理支持、风险处置 三等级", "日常交流、心理支持、风险处置三个等级"),
        ("日常交流 / 心理支持 / 风险处置 分级", "日常交流/心理支持/风险处置分级"),
        ("风险处置 报告", "风险处置报告"),
        ("识别 日常交流", "识别日常交流"),
        (" 心理支持或风险处置", "心理支持或风险处置"),
        ("以及 风险处置", "以及风险处置"),
        ("日常交流/心理支持/风险处置 分类", "日常交流/心理支持/风险处置分类"),
        ("日常交流/心理支持/风险处置 分级", "日常交流/心理支持/风险处置分级"),
    )
    count = 0
    for paragraph in paragraphs(document):
        for run in paragraph.runs:
            before = run.text
            after = before
            for old, new in fixes:
                after = after.replace(old, new)
            if after != before:
                run.text = after
                count += 1
    return count


def main() -> None:
    if STAGE_ROOT.exists():
        shutil.rmtree(STAGE_ROOT)
    STAGE_ROOT.mkdir(parents=True)
    for path in sorted(SOURCE_ROOT.glob("*.docx")):
        document = Document(path)
        changed = sum(replace_runs(paragraph) for paragraph in paragraphs(document))
        changed += post_fixes(document)
        output = STAGE_ROOT / path.name
        document.save(output)
        print(path.name, "changed_runs=", changed, "size=", output.stat().st_size)


if __name__ == "__main__":
    main()
