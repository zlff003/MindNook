from __future__ import annotations

from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document


ROOT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\risk_model_updated")


for path in sorted(ROOT.glob("*.docx")):
    with ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"{path.name}: corrupt part {bad}")
    doc = Document(path)
    content = "\n".join(p.text for p in doc.paragraphs)
    table_content = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    all_text = content + "\n" + table_content
    old_route_patterns = (
        "低风险",
        "中风险",
        "psychological_support",
        "safety_risk",
        "普通交流不进入心理评估流程",
        "心理支持与安全风险得到差异化处理",
    )
    old_hits = [value for value in old_route_patterns if value in all_text]
    if old_hits:
        raise RuntimeError(f"{path.name}: old route terms remain: {old_hits}")
    if "设计方案书.docx" == path.name:
        required = (
            "CHAT 不触发心理 RAG、报告或风险工具",
            "CONSULT 生成心理报告和台账，不创建个案",
            "RISK 生成报告、台账和个案",
            "预警任务依赖个案创建成功",
            "risk_level：CHAT、CONSULT 或 RISK",
        )
        missing = [value for value in required if value not in all_text]
        if missing:
            raise RuntimeError(f"{path.name}: missing core statements: {missing}")
    print(
        path.name,
        f"size={path.stat().st_size}",
        f"paragraphs={len(doc.paragraphs)}",
        f"tables={len(doc.tables)}",
        f"CHAT={all_text.count('CHAT')}",
        f"CONSULT={all_text.count('CONSULT')}",
        f"RISK={all_text.count('RISK')}",
    )
