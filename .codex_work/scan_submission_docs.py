from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document


ROOT = Path(r"D:\Projects\Agent\MindBridge\项目提交文档")
TERMS = re.compile(
    r"低风险|中风险|高风险|风险等级|风险分级|general|psychological_support|safety_risk|"
    r"CHAT|CONSULT|RISK|报告|台账|个案|预警"
)


def scan_docx(path: Path) -> None:
    doc = Document(path)
    print(f"\n===== {path.name} | paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} =====")
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text and TERMS.search(text):
            print(f"P{index:04d}\t{paragraph.style.name}\t{text}")
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            joined = " | ".join(values)
            if TERMS.search(joined):
                print(f"T{ti}R{ri}\t{joined}")
    with ZipFile(path) as archive:
        xml_names = [
            name
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        ]
        all_xml = "\n".join(archive.read(name).decode("utf-8", "ignore") for name in xml_names)
        print("XML_TERM_COUNTS", {term: all_xml.count(term) for term in ("低风险", "中风险", "高风险", "CHAT", "CONSULT", "RISK")})


for file_path in sorted(ROOT.glob("*.docx")):
    scan_docx(file_path)
