from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\Projects\Agent\MindBridge\项目提交文档\MindNook智能体 - 设计方案书.docx")
OUTPUT = Path(r"D:\Projects\Agent\MindBridge\MindNook\MindNook智能体 - 设计方案书.updated.docx")


def find_unique_heading(document: Document, text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one heading {text!r}, found {len(matches)}")
    return matches[0]


def add_before(document: Document, anchor: Paragraph, text: str, style: str) -> Paragraph:
    paragraph = document.add_paragraph(text, style=style)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_labeled_before(
    document: Document,
    anchor: Paragraph,
    label: str,
    body: str,
    style: str = "Normal",
) -> Paragraph:
    paragraph = document.add_paragraph(style=style)
    lead = paragraph.add_run(label)
    lead.bold = True
    paragraph.add_run(body)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def remove_between(start: Paragraph, end: Paragraph) -> None:
    current = start._p.getnext()
    while current is not None and current is not end._p:
        following = current.getnext()
        current.getparent().remove(current)
        current = following


def main() -> None:
    document = Document(SOURCE)
    chapter4 = find_unique_heading(document, "4 多智能体协作机制")
    chapter5 = find_unique_heading(document, "5 模型、上下文与知识设计")
    remove_between(chapter4, chapter5)

    add_before(
        document,
        chapter5,
        "MindNook 采用 CoordinatorAgent 统一调度、CollaborationBlackboard 共享协作的运行机制。"
        "每轮对话建立独立黑板，专业智能体不直接互相调用，而是通过任务、结构化产物、协作消息和事件交换结果。"
        "该机制将意图理解、风险控制、上下文组织、回复生成和安全复审解耦，同时保留完整的协作轨迹。",
        "Normal",
    )

    add_before(document, chapter5, "4.1 CollaborationBlackboard设计", "Heading 2")
    add_before(
        document,
        chapter5,
        "CollaborationBlackboard 是单轮 Agent Runtime 内的请求级协作空间。它不承担身份认证、数据库访问、心理报告落库或后台工具执行；这些业务职责仍由 AgentHarness 完成。黑板只负责承载本轮推理所需的协作状态，并为 CoordinatorAgent 提供统一调度依据。",
        "Normal",
    )
    add_before(document, chapter5, "其设计遵循以下原则：", "Normal")
    for item in (
        "（1）请求隔离：每轮对话使用独立 run_id 和根任务，任务及产物不跨学生、会话或请求混用。",
        "（2）能力驱动：任务声明所需 capability，由具备对应能力的智能体认领，避免角色之间形成硬编码调用链。",
        "（3）追加留痕：Task、Artifact、Message 和 Event 创建后不原地覆盖，状态变化通过追加事件表达。",
        "（4）结构化交换：智能体以确定的数据结构发布意图、风险、上下文、候选回复和审查结果，减少自由文本歧义。",
        "（5）安全优先：固定安全规则和 SAFETY_OVERRIDE 构成不可降低的风险下限；候选回复未经审查不得对学生公开。",
        "（6）受控终止：CoordinatorAgent 统一控制协作轮数、模型调用次数和总耗时，达到上限后进入确定性降级。",
    ):
        add_before(document, chapter5, item, "Normal")

    add_before(document, chapter5, "4.2 核心协作对象", "Heading 2")
    add_before(
        document,
        chapter5,
        "黑板以五类核心对象描述“要完成什么、形成了什么、谁通知了谁以及发生了什么”，并由 RuntimeTrace 记录运行预算、步骤和降级状态。",
        "Normal",
    )
    add_labeled_before(
        document,
        chapter5,
        "AgentTask：",
        "表示待处理事项，包含任务类型、所需能力、创建者、父任务、输入产物和元数据。当前任务覆盖根协调、输入理解、风险评估、上下文收集、回复准备、候选生成、安全审查和回复修复。任务对象创建后保持不可变。",
    )
    add_labeled_before(
        document,
        chapter5,
        "AgentArtifact：",
        "表示结构化工作产物，包括 hard_rule、intent、risk、context、response_prompt、response_proposal、safety_review 和 final_response。同类产物再次生成时追加新版本，旧版本继续保留。",
    )
    add_labeled_before(
        document,
        chapter5,
        "AgentMessage：",
        "表示智能体之间的协作通知，可关联任务和产物 ID，例如 ResponseAgent 通知 SafetyAgent 审查候选回复。消息只用于协作和可观测性，不触发智能体之间的直接方法调用。",
    )
    add_labeled_before(
        document,
        chapter5,
        "AgentEvent：",
        "表示黑板中已经发生的事实，包括任务创建、认领与关闭，产物发布，安全覆盖、修复请求、预算耗尽和最终采纳。事件序号单调递增，是还原协作顺序和审计判断过程的依据。",
    )
    add_labeled_before(
        document,
        chapter5,
        "AgentTurnResult：",
        "表示一个智能体完成一次任务后的统一返回，能够同时追加 Message、Artifact、Task 和 Event，并声明是否关闭当前任务。黑板完成合法性校验后一次性合并结果。",
    )

    add_before(document, chapter5, "4.3 任务认领与结果发布", "Heading 2")
    add_before(
        document,
        chapter5,
        "CoordinatorAgent 根据黑板中缺失的产物派生下一项任务，并在任务中声明 capability。各专业智能体轮询开放任务，只能认领与自身能力匹配且尚未被认领的任务。当前映射关系为：UnderstandingAgent 认领 understanding；SafetyAgent 认领 risk_assessment 和 safety_review；ContextAgent 认领 context；ResponseAgent 认领 response_prepare、response_stream 和 response_repair。",
        "Normal",
    )
    add_before(
        document,
        chapter5,
        "任务被认领后，黑板追加 TASK_CLAIMED 事件。智能体读取任务指定的输入产物完成处理，并返回 AgentTurnResult。merge_turn() 首先校验任务认领者与返回者一致，再依次追加产物、消息、后续任务和业务事件；若 close_task 为真，则最后追加 TASK_CLOSED。任务是否开放、由谁认领及是否完成，均由事件计算，而不是修改原任务对象。",
        "Normal",
    )
    add_before(
        document,
        chapter5,
        "结果发布采用版本化和单一采纳原则。同类 Artifact 可以存在多个版本，但不能覆盖历史记录；SafetyAgent 的审查结果必须关联明确的候选回复版本；只有 CoordinatorAgent 可以发布 final_response 和 FINAL_ADOPTED。若任务被重复认领、输入产物缺失、返回者不匹配或协作无法继续，Runtime 将记录异常并进入受控降级，避免产生不确定状态。",
        "Normal",
    )

    add_before(document, chapter5, "4.4 单轮对话协作流程", "Heading 2")
    add_before(document, chapter5, "一轮学生对话按照以下顺序执行：", "Normal")
    for item in (
        "（1）AgentHarness 完成学生与会话归属校验、请求幂等检查、PII 脱敏和 MemoryContext 装配，再调用 CoordinatorAgent。",
        "（2）CoordinatorAgent 创建 coordinate_turn 根任务，执行固定安全预检并发布 hard_rule；明确危险表达同时产生 SAFETY_OVERRIDE。",
        "（3）UnderstandingAgent 认领 understand_input，发布包含对话类型、主题、知识需求和理由的 intent。",
        "（4）普通交流由 CoordinatorAgent 直接形成无需模型风险评估的 risk；心理支持或安全风险则由 SafetyAgent 认领 assess_risk，并将模型结果与固定规则风险下限合并。",
        "（5）ContextAgent 认领 collect_context，装配 memoryBrief、modelHistory、回复约束和适用 Skill；心理支持及中高风险场景同时执行混合 RAG，失败时降级为内置审核知识。",
        "（6）ResponseAgent 依次完成 prepare_response 和 generate_response，生成内部 response_proposal，并发布 review_response 任务。",
        "（7）SafetyAgent 对候选回复执行发送前审查；通过后由 CoordinatorAgent 发布 final_response，未通过则进入一次受限修复。",
        "（8）AgentHarness 在同一事务中保存最终消息、Agent Run、步骤、协作事件和必要的心理报告；事务提交后异步分发台账、个案和预警任务，最后通过 SSE 返回安全成品。",
    ):
        add_before(document, chapter5, item, "Normal")
    add_before(
        document,
        chapter5,
        "三类业务路径共享上述主流程，但副作用范围不同：普通交流不触发心理 RAG、心理报告或风险工具；低风险心理支持生成报告和台账；中风险额外创建风险个案；高风险生成报告、台账和个案，并在个案创建成功后触发预警。所有后台任务均不得阻塞学生获得回复。",
        "Normal",
    )

    add_before(document, chapter5, "4.5 回复审查与修复流程", "Heading 2")
    add_before(
        document,
        chapter5,
        "ResponseAgent 生成的 Token 仅构成内部候选流。系统先在服务端累积完整候选，再由 SafetyAgent 对其执行模型审查与确定性输出规则检查，审查通过前不向学生发送正文。审查重点包括：是否包含诊断或药物建议、危险操作细节、后台风险信息、虚构校园资源、绝对保密承诺、歧视标签，以及高风险场景是否缺少现实安全行动或提出过多问题。",
        "Normal",
    )
    add_before(document, chapter5, "审查结果按以下规则处理：", "Normal")
    for item in (
        "（1）审查通过：SafetyAgent 发布 approved=true 的 safety_review；CoordinatorAgent 确认其对应当前 proposal 版本后采纳候选回复。",
        "（2）首次拒绝：SafetyAgent 发布 violations 和 revision_guidance；CoordinatorAgent 创建 repair_response，ResponseAgent 根据意见生成 response_proposal v2，并再次提交审查。",
        "（3）再次拒绝或预算耗尽：停止继续生成，由 CoordinatorAgent 采用经过审核的固定安全回退回复。",
        "（4）模型生成失败：ResponseAgent 提交带 trusted_fallback 标记的固定模板，SafetyAgent 对模板执行确定性审查。",
        "（5）安全审查服务不可用：仅普通交流候选在通过固定输出规则后可以放行；心理支持和高风险场景进入固定安全回退。",
    ):
        add_before(document, chapter5, item, "Normal")
    add_before(
        document,
        chapter5,
        "修复过程不会覆盖历史候选和审查结果，response_proposal v1/v2 与 safety_review v1/v2 均保留在黑板及 Trace 中。只有 SafetyAgent 审查完成、CoordinatorAgent 发布 FINAL_ADOPTED 且 AgentHarness 成功落库后，最终回复才会按片段通过 SSE 公开给学生，从而避免不安全内容先流出后无法撤回。",
        "Normal",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
