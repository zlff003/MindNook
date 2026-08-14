from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\Projects\Agent\MindBridge\项目提交文档\MindNook智能体 - 设计方案书 - 第5章重构版.docx")
OUTPUT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\MindNook智能体 - 设计方案书 - Harness与RAG更新版.docx")


def find_paragraph(document, exact_text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact_text}")


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_exact(document, old_text, new_text):
    paragraph = find_paragraph(document, old_text)
    set_paragraph_text(paragraph, new_text)


def clear_between(document, start_text, end_text):
    start = find_paragraph(document, start_text)
    end = find_paragraph(document, end_text)
    node = start._p.getnext()
    while node is not None and node is not end._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node
    return end


def insert_paragraph(anchor, text="", style="Normal", bold_lead=False):
    node = OxmlElement("w:p")
    anchor._p.addprevious(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = style
    if bold_lead and "：" in text:
        lead, rest = text.split("：", 1)
        run = paragraph.add_run(lead + "：")
        run.bold = True
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(anchor, items):
    for item in items:
        insert_paragraph(anchor, item, "List Bullet", bold_lead=True)


def add_steps(anchor, items):
    for item in items:
        insert_paragraph(anchor, item, "List Number", bold_lead=True)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for column_index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(widths[column_index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def format_table(table):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(10.5)
                    if row_index == 0:
                        run.bold = True


def insert_table(document, anchor, headers, rows, widths, style):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = style
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    set_table_geometry(table, widths)
    format_table(table)
    anchor._p.addprevious(table._tbl)
    return table


def rewrite_rag_section(document, table_style):
    anchor = clear_between(document, "5.5 混合 RAG 检索", "5.6 心理场景 Skill")
    insert_paragraph(
        anchor,
        "混合 RAG 为心理支持和风险处置提供可追溯的本地知识证据。ContextAgent 根据本轮对话类型、最终风险等级和主题决定是否检索，再将结果写入 ContextBundle.knowledge，供 ResponseAgent 生成回复。RAG 只负责回答“有哪些相关知识与现实资源”，不负责确定风险等级，也不能覆盖固定安全规则和 Skill 中的表达约束。",
    )

    insert_paragraph(anchor, "1. 检索触发策略", "Normal").runs[0].bold = True
    add_bullets(anchor, [
        "日常交流：不触发心理 RAG，只使用当前输入和必要的近期会话上下文。",
        "心理支持：将学生当前输入与 UnderstandingAgent 提取的主题、知识需求组合为检索 Query。",
        "风险处置：仅在最终 risk_level=RISK 时补充“安全风险、危机支持、当前安全、现实求助”等固定主题，并保证《安全风险与危机支持边界》占用一个 Top K 名额。",
    ])
    insert_paragraph(anchor, "安全主题由最终风险结果驱动，不依赖模型是否恰好输出相关关键词。该机制只增强安全知识召回，不改变 SafetyAgent 的风险结论；即使检索失败，风险处置仍需执行固定安全规则和现实求助引导。")

    insert_paragraph(anchor, "2. 知识库构建", "Normal").runs[0].bold = True
    add_steps(anchor, [
        "文档解析：读取 data/knowledge 第一层的 UTF-8 中文 Markdown，保留文档标题和章节结构。",
        "知识切块：按段落和句子边界切分，每块最多 512 字符，相邻块重叠 64 字符。",
        "来源标记：每个 Chunk 保存来源文件、章节、块序号和字符范围，保证证据可追溯。",
        "索引构建：生产默认使用 FastEmbed 加载 BAAI/bge-small-zh-v1.5，将 512 维中文向量写入 Chroma；同时以标题、章节和正文建立 BM25 索引。",
        "增量更新：知识语料签名变化时更新 Chroma、删除失效 Chunk，并重建 BM25 索引。",
    ])
    insert_paragraph(anchor, "知识内容不得包含临床诊断、药物推荐或危险方法细节；校园电话、联系人、办公时间和服务流程必须具有正式来源，并在正式使用前完成专业审核。")

    insert_paragraph(anchor, "3. 双路召回与融合", "Normal").runs[0].bold = True
    add_steps(anchor, [
        "双路召回：同一 Query 同时进入 Chroma 向量索引和 BM25 关键词索引，两路分别取得最多 12 个候选并对候选 ID 取并集。",
        "分数归一：对向量相似度和 BM25 相关度分别执行 Min-Max 归一化，未被某一路召回的候选在该路记为 0。",
        "加权融合：按照 BGE 向量 0.75、BM25 0.25 计算融合分，兼顾语义泛化与关键词稳定命中。",
        "Cross-Encoder 重排：将融合候选交给 BAAI/bge-reranker-base，模型分与融合分各占 0.5；超时、加载或推理异常时自动退回 Hybrid 排序。",
        "上下文扩展：对第一名 Chunk 合并同文档前后块并清除 overlap，随后执行来源去重，最终返回最多 3 条证据。",
    ])
    insert_table(
        document,
        anchor,
        ["参数", "默认值", "设计说明"],
        [
            ["Chunk 大小 / 重叠", "512 / 64 字符", "控制检索粒度并减轻语义断裂"],
            ["Embedding", "BAAI/bge-small-zh-v1.5", "FastEmbed 本地运行，生产默认 512 维"],
            ["候选数量", "每路最多 12 条", "为融合和重排提供候选范围"],
            ["融合权重", "向量 0.75 / BM25 0.25", "向量主召回，BM25 补充词面信号"],
            ["Reranker", "BAAI/bge-reranker-base", "Cross-Encoder 模型分与融合分各占 0.5"],
            ["重排超时", "30 秒", "超时或异常时回退 Hybrid 原序"],
            ["最终结果", "最多 3 条", "完成安全名额、邻块扩展和来源去重后注入"],
        ],
        [2300, 2900, 3800],
        table_style,
    )

    insert_paragraph(anchor, "4. 注入与降级", "Normal").runs[0].bold = True
    add_bullets(anchor, [
        "证据注入：最终结果连同来源文件、章节、融合分、重排分和扩展信息写入 ContextBundle.knowledge。",
        "事实边界：检索文本只作为事实参考，不得作为系统指令执行；具体电话、人员、时间和政策必须在证据中明确存在。",
        "职责边界：RAG 不能改变风险等级或安全下限；RAG 提供知识与资源，Skill 负责场景表达。",
        "分级降级：Cross-Encoder 不可用时退回 Hybrid 排序；Chroma、BM25 或知识文件异常时使用少量内置审核知识，并记录 rag_fallback。",
        "五阶段观测：保留 vector、bm25、hybrid、reranked 和 final 结果，业务使用 final，评测通过五阶段比较定位召回、融合、重排和扩展问题。",
    ])


def rewrite_design_testing(document):
    anchor = clear_between(document, "8.1 异常与降级", "8.2 测试与验收")
    insert_paragraph(anchor, "系统按故障层级执行确定性降级：模型失败使用固定模板；Redis失败从MySQL回填；Cross-Encoder加载、推理或超时时退回Hybrid排序；Chroma、BM25或知识文件异常时使用内置审核知识；工具失败进入重试与死信；关键Skill缺失时使用固定安全回复。所有降级均写入Step与Trace，不得降低固定风险下限或绕过回复安全审查。")

    anchor = clear_between(document, "8.2 测试与验收", "8.3 部署与当前限制")
    insert_paragraph(anchor, "系统采用“单元与集成测试—Engineering Harness工程回归—独立RAG严格评测—真实环境与人工评审”四层验收方式。")
    add_bullets(anchor, [
        "单元与集成测试：覆盖多智能体协议、记忆隔离、风险硬规则、Skill路由、RAG链路和工具队列。",
        "Engineering Harness：提供risk-safety、agent-routing、standard-skills、rag、api和tool-queue六个固定套件；每个套件使用临时SQLite、内存短期记忆、确定性Mock模型、临时RAG索引和临时Excel，不污染正式环境。",
        "独立RAG严格评测：使用108条样例比较vector、bm25、hybrid、reranked和final五个阶段，并检查Hit@1、Recall@3、安全召回、误触发率和上下文重复率。",
        "上线验收：补充真实模型API、MySQL、Redis、异步Worker、HTTP/SSE、浏览器UI、性能稳定性和心理专业人工盲审。",
    ])
    insert_paragraph(anchor, "2026年8月14日实际运行Engineering Harness，六个套件全部通过，共21项检查全部通过。工程回归通过仅表示确定性链路和最低指标未退化，不等同于严格RAG发布门禁或真实模型回答质量已经达标。")

    anchor = clear_between(document, "8.3 部署与当前限制", "第三部分 测试方案与结果")
    insert_paragraph(anchor, "部署除模型、MySQL、Redis、Chroma、Skill目录、执行预算和工具Worker外，还需配置FastEmbed模型缓存。BAAI/bge-small-zh-v1.5首次使用约需下载90MB，BAAI/bge-reranker-base约需下载1GB；离线部署应提前预热data/models，Cross-Encoder会增加CPU推理耗时，并通过30秒超时回退Hybrid排序。")
    insert_paragraph(anchor, "当前严格RAG验收仍未全部通过：Hit@1为0.888，Recall@3为0.925，后者低于0.950门槛；安全Hit@1和Recall@3均为1.000，普通问题误触发率和上下文重复率均为0。完整Runtime的RAG触发召回率为0.562，对隐式心理支持、长上下文、功能影响和混合意图仍需优化。现有108条样例仍是单一回归集，后续需要建立不参与调参的保留集，并加入经脱敏、审核的真实表达。")
    insert_paragraph(anchor, "Engineering Harness使用临时基础设施、Mock模型和进程内工具执行器，不覆盖真实模型回答质量、真实MySQL/Redis/Worker、浏览器UI、并发压测、故障注入和长时间稳定性；这些能力必须在上线前单独验证。")


def update_test_part(document, table_style):
    anchor = clear_between(document, "2.2 模型、知识库与工具配置", "2.3 测试数据准备")
    insert_paragraph(anchor, "自动化测试将确定性工程回归与正式RAG质量评测分开配置，避免模型网络波动影响基础链路判断。")
    insert_table(
        document,
        anchor,
        ["测试对象", "配置", "用途"],
        [
            ["模型服务", "Engineering Harness固定使用MockModelProvider", "离线验证路由、生成、审查和持久化协议"],
            ["数据库", "每个套件使用全新临时SQLite", "隔离业务数据，验证消息、报告、Trace和ToolJob"],
            ["短期记忆", "内存Redis兼容实现", "验证记忆读取、写回和会话隔离"],
            ["工程RAG", "临时索引、Hash测试向量、关闭Cross-Encoder", "稳定验证108条数据集覆盖和工程指标底线"],
            ["严格RAG", "BGE 512维Embedding + BGE Cross-Encoder", "验证正式检索配置与发布级指标"],
            ["工具输出", "临时Excel与进程内执行器", "验证ledger/case/alert DAG而不写正式台账"],
        ],
        [1900, 3900, 3200],
        table_style,
    )

    anchor = clear_between(document, "3.5 自动化测试", "4 功能与业务场景测试")
    insert_paragraph(anchor, "自动化验证分为三步：先运行pytest单元与集成测试，再执行一键Engineering Harness，最后运行正式RAG严格评测。Harness支持选择单个或多个套件，并可输出机器可读JSON；任一套件失败时进程返回非零退出码，可直接接入CI门禁。")
    add_bullets(anchor, [
        "工程回归命令：python -m app.harness.runner。",
        "机器报告命令：python -m app.harness.runner --json。",
        "RAG发布门禁：python -m app.rag.evaluation --evaluate-answers --strict。",
        "真实模型与人工评审：另行执行Provider冒烟、批量回答质量评测和心理专业人工盲审。",
    ])

    anchor = clear_between(document, "4.4 RAG 与知识使用测试", "4.5 工具调用与任务执行测试")
    insert_paragraph(anchor, "RAG测试覆盖触发路由、索引构建、双路召回、Cross-Encoder重排、安全知识兜底、证据注入和异常降级。")
    add_bullets(anchor, [
        "触发策略：日常交流不触发心理RAG；心理支持触发主题检索；风险处置补充安全主题并固定安全知识Top K名额。",
        "索引与召回：验证512/64切块、BGE向量、BM25候选、0.75/0.25融合及来源元数据。",
        "重排与扩展：比较不重排、旧词面公式和BGE Cross-Encoder，验证首块前后文扩展和来源去重。",
        "安全与降级：验证Reranker超时退回Hybrid、索引异常退回内置知识，并确认降级不降低风险安全下限。",
        "质量指标：统计Hit@1、Recall@3、MRR@3、nDCG@3、Top1证据召回、安全召回、路由误触发率和上下文重复率。",
    ])

    anchor = clear_between(document, "8.1 测试执行概况", "8.2 功能测试结果")
    insert_paragraph(anchor, "2026年8月14日执行当前版本Engineering Harness，结果为PASS：6个套件全部通过，21项检查全部通过。独立RAG评测使用108条样例，其中89条检索样例、19条普通问题负样例、13条安全关键样例和9条跨Chunk证据样例。")

    anchor = clear_between(document, "8.2 功能测试结果", "8.3 非功能测试结果")
    insert_table(
        document,
        anchor,
        ["Harness套件", "结果", "主要验证内容"],
        [
            ["risk-safety", "通过", "高风险识别、报告、元数据隔离和工具入队"],
            ["agent-routing", "通过", "三级路由、Agent步骤和协作记录"],
            ["standard-skills", "通过", "必需Skill、场景选择和Workflow注入"],
            ["rag", "通过工程底线", "108条数据集、召回指标、安全召回和重复率"],
            ["api", "通过", "健康检查、认证授权、公开SSE和知识库管理接口"],
            ["tool-queue", "通过", "持久化DAG、幂等、执行、限流和死信"],
        ],
        [2500, 1800, 4700],
        table_style,
    )
    insert_paragraph(anchor, "Harness的RAG套件使用稳定工程回归底线，rag套件通过不代表正式BGE与Cross-Encoder配置已经满足严格发布阈值。")

    anchor = clear_between(document, "8.6 指标对比与验收结论", "9 测试结论与改进建议")
    insert_table(
        document,
        anchor,
        ["严格RAG指标", "实际值", "要求", "结论"],
        [
            ["Hit@1", "0.888", ">= 0.850", "通过"],
            ["Recall@3", "0.925", ">= 0.950", "未通过"],
            ["安全Hit@1", "1.000", "= 1.000", "通过"],
            ["安全Recall@3", "1.000", "= 1.000", "通过"],
            ["普通问题误触发率", "0.000", "<= 0.020", "通过"],
            ["上下文重复率", "0.000", "<= 0.150", "通过"],
        ],
        [3300, 1500, 2100, 2100],
        table_style,
    )
    insert_paragraph(anchor, "辅助指标为MRR@3=0.919、nDCG@3=0.905、Top1证据召回=1.000。综合结论：Engineering Harness工程门禁通过，但严格RAG发布验收因Recall@3未达到0.950而未通过。")

    anchor = clear_between(document, "9.2 风险与限制", "9.3 上线建议")
    add_bullets(anchor, [
        "检索覆盖风险：Recall@3仍差2.5个百分点，混合意图、长上下文、功能影响和边界问题仍是主要失败类型。",
        "路由遗漏风险：完整Runtime的RAG触发召回率为0.562，部分隐式心理支持表达未触发检索。",
        "评测过拟合风险：108条样例仍为单一回归集，需要建立独立保留集并引入经审核的真实表达。",
        "部署资源风险：Embedding和Cross-Encoder需要预下载模型并增加本地存储、内存和CPU推理开销。",
        "验证边界：Engineering Harness不覆盖真实模型回答质量、真实基础设施、浏览器UI、并发压测和长稳运行。",
    ])


def main():
    document = Document(SOURCE)
    table_style = document.tables[0].style if document.tables else "Table Grid"

    replace_exact(
        document,
        "本地混合RAG负责为心理支持和风险场景提供知识依据。系统对知识库文件进行切分，同时执行Chroma向量检索和BM25关键词检索，再经过分数融合、本地重排和相邻上下文扩展，将最终检索内容交给ContextAgent。RAG异常时系统可降级使用内置知识，不阻断基础对话和风险回复。",
        "本地混合RAG负责为心理支持和风险处置提供知识依据。生产配置使用BAAI/bge-small-zh-v1.5生成512维中文向量，并结合BM25完成双路召回；候选按0.75/0.25融合后，由BAAI/bge-reranker-base执行Cross-Encoder重排，再进行安全知识固定名额、首块前后文扩展和来源去重。Reranker异常时回退Hybrid排序，索引异常时使用内置审核知识，不阻断基础对话和风险安全回复。",
    )
    replace_exact(
        document,
        "知识检索采用 Chroma 向量召回与 BM25 关键词召回的混合方式，经过分数融合、本地重排和相邻上下文扩展后返回相关证据。ContextAgent 另由 SkillRegistry 按等级与主题选择心理场景 Skill：日常交流不加载；心理支持加载基础支持、转介及匹配的焦虑、睡眠或学业 Skill；风险处置加载高风险安全计划与转介 Skill。",
        "知识检索采用Chroma中的BGE中文向量召回与BM25关键词召回，按0.75/0.25融合后使用BGE Cross-Encoder重排，并执行安全知识固定名额、相邻上下文扩展和来源去重。ContextAgent另由SkillRegistry按等级与主题选择心理场景Skill：日常交流不加载；心理支持加载基础支持、转介及匹配的焦虑、睡眠或学业Skill；风险处置加载高风险安全计划与转介Skill。",
    )

    rewrite_rag_section(document, table_style)
    rewrite_design_testing(document)
    update_test_part(document, table_style)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    all_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    required = [
        "BAAI/bge-small-zh-v1.5",
        "BAAI/bge-reranker-base",
        "向量 0.75 / BM25 0.25",
        "6个套件全部通过，21项检查全部通过",
        "Recall@3为0.925",
        "完整Runtime的RAG触发召回率为0.562",
    ]
    missing = [item for item in required if item not in all_text and not any(item in cell.text for table in reopened.tables for row in table.rows for cell in row.cells)]
    if missing:
        raise AssertionError(f"Missing expected content: {missing}")
    stale = ["向量 0.45、BM25 0.55", "0.65 / 0.15 / 0.12 / 0.08", "本地RAG的严格质量基线尚未全部达标"]
    stale_found = [item for item in stale if item in all_text or any(item in cell.text for table in reopened.tables for row in table.rows for cell in row.cells)]
    if stale_found:
        raise AssertionError(f"Stale content remains: {stale_found}")
    with ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise AssertionError("Invalid DOCX package")
    print(f"output={OUTPUT}")
    print(f"size={OUTPUT.stat().st_size}")
    print(f"paragraphs={len(reopened.paragraphs)} tables={len(reopened.tables)} images={len(reopened.inline_shapes)}")


if __name__ == "__main__":
    main()
