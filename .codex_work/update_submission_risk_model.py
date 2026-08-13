from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import shutil

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE_ROOT = Path(r"D:\Projects\Agent\MindBridge\项目提交文档")
STAGE_ROOT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\risk_model_updated")


@dataclass
class Change:
    file: str
    location: str
    before: str
    after: str


CHANGES: list[Change] = []


COMMON_REPLACEMENTS = (
    ("general / psychological_support / safety_risk", "CHAT / CONSULT / RISK"),
    ("general、psychological_support、safety_risk", "CHAT、CONSULT、RISK"),
    ("psychological_support", "CONSULT"),
    ("safety_risk", "RISK"),
    ("general", "CHAT"),
    ("中、高风险", "RISK"),
    ("中高风险", "RISK"),
    ("低风险", "CONSULT"),
    ("中风险", "RISK"),
    ("高风险", "RISK"),
    ("普通交流", "CHAT"),
    ("心理支持", "CONSULT"),
    ("安全风险", "RISK"),
)


def replace_runs(paragraph: Paragraph, replacements=COMMON_REPLACEMENTS) -> bool:
    changed = False
    for run in paragraph.runs:
        value = run.text
        updated = value
        for old, new in replacements:
            updated = updated.replace(old, new)
        if updated != value:
            run.text = updated
            changed = True
    return changed


def set_paragraph(paragraph: Paragraph, text: str) -> None:
    if len(paragraph.runs) == 1:
        paragraph.runs[0].text = text
    else:
        paragraph.text = text


def paragraph_index(document: Document, paragraph: Paragraph) -> int:
    return next(i for i, candidate in enumerate(document.paragraphs) if candidate._p is paragraph._p)


def all_document_paragraphs(document: Document):
    seen: set[int] = set()
    for paragraph in document.paragraphs:
        if id(paragraph._p) not in seen:
            seen.add(id(paragraph._p))
            yield paragraph, "P"
    for ti, table in enumerate(document.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, paragraph in enumerate(cell.paragraphs):
                    if id(paragraph._p) not in seen:
                        seen.add(id(paragraph._p))
                        yield paragraph, f"T{ti}R{ri}C{ci}P{pi}"
    for si, section in enumerate(document.sections):
        for region_name, region in (("H", section.header), ("F", section.footer)):
            for pi, paragraph in enumerate(region.paragraphs):
                if id(paragraph._p) not in seen:
                    seen.add(id(paragraph._p))
                    yield paragraph, f"S{si}{region_name}P{pi}"
            for ti, table in enumerate(region.tables):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, paragraph in enumerate(cell.paragraphs):
                            if id(paragraph._p) not in seen:
                                seen.add(id(paragraph._p))
                                yield paragraph, f"S{si}{region_name}T{ti}R{ri}C{ci}P{pi}"


def apply_common(document: Document, file_name: str) -> None:
    paragraph_index = {id(p._p): i for i, p in enumerate(document.paragraphs)}
    for paragraph, location in all_document_paragraphs(document):
        before = paragraph.text
        if replace_runs(paragraph):
            after = paragraph.text
            loc = f"P{paragraph_index[id(paragraph._p)]:04d}" if id(paragraph._p) in paragraph_index else location
            CHANGES.append(Change(file_name, loc, before, after))


def override(document: Document, file_name: str, startswith: str, new_text: str) -> None:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"{file_name}: expected one paragraph starting {startswith!r}, got {len(matches)}")
    paragraph = matches[0]
    before = paragraph.text
    set_paragraph(paragraph, new_text)
    index = paragraph_index(document, paragraph)
    CHANGES.append(Change(file_name, f"P{index:04d}", before, new_text))


def override_contains(document: Document, file_name: str, needle: str, new_text: str) -> None:
    matches = [p for p in document.paragraphs if needle in p.text]
    if len(matches) != 1:
        raise RuntimeError(f"{file_name}: expected one paragraph containing {needle!r}, got {len(matches)}")
    paragraph = matches[0]
    before = paragraph.text
    set_paragraph(paragraph, new_text)
    index = paragraph_index(document, paragraph)
    CHANGES.append(Change(file_name, f"P{index:04d}", before, new_text))


def remove_startswith(document: Document, file_name: str, startswith: str) -> None:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"{file_name}: expected one paragraph to remove starting {startswith!r}, got {len(matches)}")
    paragraph = matches[0]
    index = paragraph_index(document, paragraph)
    CHANGES.append(Change(file_name, f"P{index:04d}", paragraph.text, "[删除：与新的三等级模型重复]"))
    paragraph._element.getparent().remove(paragraph._element)


def update_main(document: Document, file_name: str) -> None:
    # Core business rule and scenario list.
    override_contains(
        document,
        file_name,
        "对话异步形成内部报告与工作台账",
        "系统按 CHAT、CONSULT、RISK 三个等级处理：CHAT 不触发心理 RAG、报告或风险工具；CONSULT 生成心理报告和台账，不创建个案；RISK 生成报告、台账和个案，预警任务依赖个案创建成功。",
    )
    override(
        document,
        file_name,
        "阶段五——人工跟进",
        "阶段五——人工跟进：RISK 信息进入授权工作人员工作台，由人员接手、联系现实支持、记录后续安排并完成复访。",
    )
    override(
        document,
        file_name,
        "1、CHAT",
        "1、CHAT：学生询问学习、编程、作业、校园事务或日常问题，系统直接回答，不触发心理 RAG、报告或风险工具。",
    )
    override(
        document,
        file_name,
        "2、一般CONSULT",
        "2、CONSULT：学生表达压力、焦虑、低落、失眠、关系或适应困扰，系统提供温和回应和少量行动建议，生成心理报告和台账，但不创建个案。",
    )
    override(
        document,
        file_name,
        "3、RISK持续困扰",
        "3、RISK：出现需要人工跟进的风险表达时，系统生成报告、台账和唯一风险个案；预警任务仅在个案创建成功后执行。",
    )
    remove_startswith(document, file_name, "4、RISK安全响应")
    override(document, file_name, "5、工作人员个案跟进", "4、工作人员个案跟进：授权人员查看必要信息、确认接手、指定负责人、添加备注、安排复访并更新状态。")
    override(document, file_name, "6、知识与资源维护", "5、知识与资源维护：管理员上传、更新、停用或删除知识来源，维护校园资源及其来源、更新时间和有效状态。")
    override(document, file_name, "7、运行异常处理", "6、运行异常处理：运营人员查看失败任务、通知状态和必要审计，执行重试或人工恢复，但无权默认读取个案正文。")

    # Use cases, tools, admin rules and consistency.
    override(
        document,
        file_name,
        "2、RISK安全响应",
        "2、RISK 安全响应：固定规则命中明确危险表达后，系统生成安全优先候选回复并完成安全检查。学生先获得安全回复；报告、台账和个案异步执行，预警任务等待个案创建成功后再执行。",
    )
    override(
        document,
        file_name,
        "1．报告、台账、个案、预警等后台任务",
        "1．CONSULT 的报告、台账任务以及 RISK 的报告、台账、个案、预警任务均应异步执行，不得阻塞学生回复。",
    )
    override(
        document,
        file_name,
        "3．RISK报告自动且幂等创建个案",
        "3．仅 RISK 报告自动且幂等创建个案；CHAT 不创建业务记录，CONSULT 只生成报告和台账，不创建个案。",
    )
    override(
        document,
        file_name,
        "5．RISK预警发送给配置的授权接收人",
        "5．RISK 预警发送给配置的授权接收人，且预警任务必须依赖个案创建成功；发送失败后自动重试、限频去重，多次失败转人工异常。",
    )
    override(
        document,
        file_name,
        "一致性：风险等级与报告",
        "一致性：CHAT、CONSULT、RISK 与 RAG、报告、台账、个案和预警动作符合统一业务规则。",
    )
    override(
        document,
        file_name,
        "唯一性：同一报告不重复创建成功台账",
        "唯一性：同一报告不重复创建成功台账或 RISK 个案；同一逻辑预警受个案依赖、限频和去重控制。",
    )
    override(
        document,
        file_name,
        "P0——安全闭环最小可用版本",
        "P0——安全闭环最小可用版本：学生登录与对话、CHAT/CONSULT/RISK 三等级分流、固定安全规则、RISK 安全检查和回退、按等级生成报告/台账/个案/预警的异步幂等闭环、最小个案工作台、权限隔离、敏感信息保护、审计及经审核的校园资源。",
    )
    override(
        document,
        file_name,
        "学校RISK判定",
        "学校 RISK 判定、个案接手、转交、关闭、复访和人工跟进时限。",
    )

    # Acceptance and traceability.
    override(
        document,
        file_name,
        "AC-01 CHAT",
        "AC-01 CHAT：学生询问普通问题时系统自然回答，不进行不必要的心理评估，不触发心理 RAG、报告或风险工具。",
    )
    override(
        document,
        file_name,
        "AC-02 CONSULT",
        "AC-02 CONSULT：系统能识别压力、焦虑、低落、睡眠和关系困扰，回答先回应处境再提供少量具体建议，并生成心理报告和台账，不创建个案。",
    )
    override(
        document,
        file_name,
        "AC-03 RISK",
        "AC-03 RISK：系统生成报告、台账和唯一风险个案；预警任务依赖个案创建成功，学生端不显示等级或内部信息。",
    )
    override(
        document,
        file_name,
        "AC-04 RISK",
        "AC-04 RISK 安全响应：明确危险表达被固定规则识别，回复优先当前安全、现实支持和短期行动，不含危险细节且经过安全检查。",
    )
    override(
        document,
        file_name,
        "AC-05 RISK闭环",
        "AC-05 RISK 闭环：系统生成报告、台账和个案，个案创建成功后执行预警；预警失败不影响学生回复，并在多次失败后进入人工异常列表。",
    )
    override(
        document,
        file_name,
        "TR-02 CONSULT",
        "TR-02 CONSULT：PRD 4.4、6.1、11.2 → OBJ-03 → FR-DLG-05、FR-RPT/ADM 相关能力 → AC-02 → CONSULT 回复、报告与台账链路测试。",
    )
    override(
        document,
        file_name,
        "TR-03 RISK安全",
        "TR-03 RISK 安全：PRD 3.2、4.4、6.2/6.3、8.4、11.3 → OBJ-04/05 → FR-SAFE-01～08、FR-TOOL-01～06、FR-ADM-03/05 → AC-03/04/05 → 固定规则、安全回复、个案依赖预警和故障演练。",
    )
    override(
        document,
        file_name,
        "TR-06 个案与幂等",
        "TR-06 个案与幂等：PRD 5.3、8.2/8.3、11.5 → OBJ-05 → FR-ADM-03/04、FR-TOOL-02 → AC-03/08 → RISK 个案唯一性、预警依赖、状态历史和并发任务测试。",
    )

    # Architecture, role model and runtime descriptions.
    override_contains(
        document,
        file_name,
        "根据风险等级生成台账、个案和预警工具计划",
        "根据 CHAT、CONSULT、RISK 等级生成后台计划：CHAT 无心理业务副作用；CONSULT 生成报告和台账；RISK 生成报告、台账和个案，并创建依赖个案成功的预警任务；",
    )
    override(
        document,
        file_name,
        "系统首先完成意图识别和风险评估",
        "系统首先完成等级识别和安全评估，再进入 CHAT、CONSULT 或 RISK 模式。CHAT 不调用心理 RAG、不生成报告或风险工具；CONSULT 调用本地心理知识，生成心理报告和台账，不创建个案；RISK 优先执行固定安全规则，生成报告、台账和风险个案，并创建依赖个案成功的预警任务。候选回复必须经过 SafetyAgent 复审，审查通过并成功落库后才能通过 SSE 返回学生端。",
    )
    override(document, file_name, "（1）CHAT：", "（1）CHAT：学习、编程、作业、校园事务和其他普通问答，不触发心理 RAG、报告或风险工具；")
    override(document, file_name, "（2）CONSULT：", "（2）CONSULT：压力、焦虑、低落、失眠、人际关系、家庭关系和校园适应等困扰，生成心理报告和台账但不创建个案；")
    override(document, file_name, "（3）RISK：", "（3）RISK：自伤、自杀、伤人、严重绝望或即时危险等需要安全响应和人工跟进的表达，生成报告、台账和个案，预警依赖个案创建成功。")
    override(
        document,
        file_name,
        "risk_level：CONSULT、RISK或RISK",
        "risk_level：CHAT、CONSULT 或 RISK；",
    )
    override(
        document,
        file_name,
        "固定安全规则与模型风险结果",
        "最终等级由业务识别结果与固定安全规则共同确定。固定规则一旦命中明确危险表达，最终等级必须为 RISK，模型不得降级为 CHAT 或 CONSULT。SafetyAgent 随即发布 SAFETY_OVERRIDE，强制进入安全优先路径。",
    )
    override(
        document,
        file_name,
        "ContextAgent 根据对话类型和风险等级决定是否触发 RAG",
        "ContextAgent 根据 CHAT、CONSULT、RISK 等级决定是否触发 RAG：",
    )
    override(document, file_name, "CHAT：只装配", "CHAT：只装配必要的近期上下文，不触发心理 RAG；")
    override(document, file_name, "CONSULT：检索", "CONSULT：检索与当前困扰相关的心理支持知识；")
    override(document, file_name, "RISK或RISK：检索", "RISK：检索心理知识，并补充“当前安全、危机支持、现实求助”等固定安全主题。")
    override(
        document,
        file_name,
        "（5）ContextAgent 认领 collect_context",
        "（5）ContextAgent 认领 collect_context，装配 memoryBrief、modelHistory、回复约束和适用 Skill；CONSULT 与 RISK 执行混合 RAG，失败时降级为内置审核知识，CHAT 不触发心理 RAG。",
    )
    override(
        document,
        file_name,
        "（8）AgentHarness 在同一事务中保存最终消息",
        "（8）AgentHarness 在同一事务中保存最终消息、Agent Run、步骤和协作事件，并按等级生成业务副作用：CHAT 不生成心理业务记录；CONSULT 保存心理报告并创建台账任务；RISK 保存报告并创建台账、个案及依赖个案成功的预警任务。事务提交后异步分发工具任务，最后通过 SSE 返回安全成品。",
    )
    override(
        document,
        file_name,
        "三类业务路径共享上述主流程",
        "三类等级共享上述主流程，但副作用严格区分：CHAT 不触发心理 RAG、报告或风险工具；CONSULT 生成心理报告和台账，不创建个案；RISK 生成报告、台账和个案，预警任务依赖个案创建成功。所有后台任务均不得阻塞学生获得回复。",
    )

    # Restore a generic security phrase that is not the conversation class.
    for paragraph in document.paragraphs:
        if "其他RISK操作" in paragraph.text:
            before = paragraph.text
            replace_runs(paragraph, (("其他RISK操作", "其他高风险操作"),))
            CHANGES.append(Change(file_name, "generic-security", before, paragraph.text))

    # Use natural Chinese compounds instead of mechanically concatenated route labels.
    compound_replacements = (
        ("CONSULT报告", "心理报告"),
        ("RISK预警", "RISK 预警"),
        ("RISK事件", "RISK 事件"),
        ("RISK任务", "RISK 任务"),
        ("RISK场景", "RISK 场景"),
        ("RISK回复", "RISK 回复"),
        ("RISK候选", "RISK 候选"),
        ("RISK识别", "RISK 识别"),
        ("RISK安全", "RISK 安全"),
        ("RISK覆盖", "RISK 覆盖"),
        ("RISK处理", "RISK 处理"),
        ("RISK跟进", "RISK 跟进"),
        ("CONSULT场景", "CONSULT 场景"),
        ("CONSULT回答", "CONSULT 回答"),
        ("CHAT候选", "CHAT 候选"),
        ("CHAT不得", "CHAT 不得"),
    )
    for paragraph in document.paragraphs:
        before = paragraph.text
        if replace_runs(paragraph, compound_replacements):
            CHANGES.append(Change(file_name, "compound-spacing", before, paragraph.text))

    # Correct terms that are identifiers or general Chinese concepts rather than route labels.
    semantic_fixes = (
        ("CHAT_support", "general_support"),
        ("获得CONSULT", "获得心理支持"),
        ("需要关注的CONSULT或风险信息", "需要关注的心理困扰或风险信息"),
        ("进行CHAT和多轮对话", "进行日常交流和多轮对话"),
        ("提供普通回答、CONSULT或安全优先回复", "提供 CHAT 普通回答、CONSULT 支持性回答或 RISK 安全优先回复"),
        ("CONSULT和风险消息形成内部评估", "CONSULT 与 RISK 消息形成内部评估"),
        ("为CONSULT与风险对话生成心理报告", "为 CONSULT 与 RISK 对话生成心理报告"),
        ("其他CONSULT场景", "CONSULT 场景"),
        ("按风险等级", "按 CHAT、CONSULT、RISK 等级"),
        ("复核风险等级", "复核 CHAT、CONSULT、RISK 等级"),
        ("包含情绪概况、风险等级", "包含情绪概况、CHAT/CONSULT/RISK 等级"),
        ("包括分类、风险等级", "包括 CHAT/CONSULT/RISK 分类等级"),
        ("覆盖风险等级与安全边界", "覆盖 CHAT/CONSULT/RISK 分级与安全边界"),
        ("最终风险等级", "最终 CHAT/CONSULT/RISK 等级"),
    )
    for paragraph in document.paragraphs:
        before = paragraph.text
        if replace_runs(paragraph, semantic_fixes):
            CHANGES.append(Change(file_name, "semantic-fix", before, paragraph.text))

    override(
        document,
        file_name,
        "10.1 风险分级",
        "10.1 CHAT / CONSULT / RISK 分级",
    )

    # Role responsibility matrix.
    table = document.tables[1]
    cell = table.rows[2].cells[4]
    before = cell.text
    cell.text = "识别 CHAT、CONSULT 或 RISK；\n提取主题、知识需求和判断理由"
    CHANGES.append(Change(file_name, "T1R2C4", before, cell.text))

    # First-part objectives table must use the same route names.
    objectives = document.tables[0]
    objective_updates = {
        (2, 2): "CHAT 不触发心理 RAG、报告或风险工具；CONSULT 与 RISK 得到差异化处理",
        (3, 1): "提供可执行的 CONSULT 支持",
        (3, 2): "CONSULT 回答先回应具体处境，再提供 1～2 个相关小步骤，并生成报告和台账",
        (4, 1): "优先处理 RISK",
        (4, 2): "明确危险信号由固定规则识别，学生及时获得安全优先回复",
        (5, 2): "RISK 生成报告、台账和个案；预警任务依赖个案创建成功，支持授权人员跟进",
    }
    for (ri, ci), value in objective_updates.items():
        cell = objectives.rows[ri].cells[ci]
        before = cell.text
        cell.text = value
        CHANGES.append(Change(file_name, f"T0R{ri}C{ci}", before, value))


def update_outline(document: Document, file_name: str) -> None:
    override(document, file_name, "6.4 RISK流程", "6.4 RISK 流程")
    override(document, file_name, "10.1 风险分级", "10.1 CHAT / CONSULT / RISK 分级")
    table = document.tables[0]
    updates = {
        (2, 2): "CHAT 不触发心理 RAG、报告或风险工具；CONSULT 与 RISK 得到差异化处理",
        (3, 1): "提供可执行的 CONSULT 支持",
        (3, 2): "CONSULT 回答先回应具体处境，再提供 1～2 个相关小步骤，并生成报告和台账",
        (4, 1): "优先处理 RISK",
        (4, 2): "明确危险信号由固定规则识别，学生及时获得安全优先回复",
        (5, 2): "RISK 生成报告、台账和个案；预警任务依赖个案创建成功，支持授权人员跟进",
    }
    for (ri, ci), value in updates.items():
        cell = table.rows[ri].cells[ci]
        before = cell.text
        cell.text = value
        CHANGES.append(Change(file_name, f"T0R{ri}C{ci}", before, value))


def update_summary(document: Document, file_name: str) -> None:
    override(
        document,
        file_name,
        "风险闭环管理",
        "风险闭环管理：系统采用 CHAT、CONSULT、RISK 三等级。CHAT 不触发心理 RAG、报告或风险工具；CONSULT 自动生成心理报告和台账，不创建个案；RISK 生成报告、台账和个案，预警任务依赖个案创建成功。管理端可查看报告、必要上下文、个案进展及运行轨迹。",
    )
    override(
        document,
        file_name,
        "形成一套可本地运行的校园心理关怀 Web 系统",
        "形成一套可本地运行的校园心理关怀 Web 系统，包括学生端、管理端、多智能体运行时、分层记忆、混合 RAG，以及按 CHAT、CONSULT、RISK 驱动的报告、台账、个案和预警闭环；同步交付项目文档、自动化测试、RAG 评测集、运行 Trace 和部署配置。预期实现三类典型对话的稳定分流，RISK 场景优先输出安全引导，关键处理过程可审计、可回放，知识和模型能力可持续扩展。",
    )


def process(path: Path) -> Path:
    document = Document(path)
    apply_common(document, path.name)
    if path.name == "MindNook智能体 - 设计方案书.docx":
        update_main(document, path.name)
    elif path.name == "MindNook智能体 - 设计方案书 - 大纲.docx":
        update_outline(document, path.name)
    elif path.name == "MindNook智能体 - 项目概要介绍.docx":
        update_summary(document, path.name)
    output = STAGE_ROOT / path.name
    document.save(output)
    return output


def main() -> None:
    if STAGE_ROOT.exists():
        shutil.rmtree(STAGE_ROOT)
    STAGE_ROOT.mkdir(parents=True)
    outputs = [process(path) for path in sorted(SOURCE_ROOT.glob("*.docx"))]
    log = STAGE_ROOT / "change_log.tsv"
    with log.open("w", encoding="utf-8") as handle:
        handle.write("file\tlocation\tbefore\tafter\n")
        for change in CHANGES:
            handle.write(
                f"{change.file}\t{change.location}\t"
                f"{change.before.replace(chr(9), ' ').replace(chr(10), ' / ')}\t"
                f"{change.after.replace(chr(9), ' ').replace(chr(10), ' / ')}\n"
            )
    for output in outputs:
        print(output)
    print(f"changes={len(CHANGES)}")
    print(log)


if __name__ == "__main__":
    main()
