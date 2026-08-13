from __future__ import annotations

from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document


ROOT = Path(r"D:\Projects\Agent\MindBridge\MindNook\.codex_work\risk_model_updated")
OLD = re.compile(r"低风险|中风险|高风险|普通交流|心理支持|安全风险|general|psychological_support|safety_risk")
SUSPICIOUS = re.compile(
    r"CONSULT.{0,35}(个案|预警)|"
    r"(个案|预警).{0,35}CONSULT|"
    r"预警.{0,35}(依赖|个案|成功)|"
    r"RISK操作|RISK账号|RISK权限|RISK任务"
)


for path in sorted(ROOT.glob("*.docx")):
    doc = Document(path)
    print(f"\n===== {path.name} =====")
    old_hits = []
    suspicious = []
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if OLD.search(text):
            old_hits.append((f"P{i:04d}", text))
        if SUSPICIOUS.search(text):
            suspicious.append((f"P{i:04d}", text))
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            text = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
            if OLD.search(text):
                old_hits.append((f"T{ti}R{ri}", text))
            if SUSPICIOUS.search(text):
                suspicious.append((f"T{ti}R{ri}", text))
    print("OLD_HITS", len(old_hits))
    for item in old_hits:
        print("OLD", *item, sep="\t")
    print("SUSPICIOUS", len(suspicious))
    for item in suspicious:
        print("CHECK", *item, sep="\t")
    print(
        "COUNTS",
        {
            "CHAT": sum(p.text.count("CHAT") for p in doc.paragraphs),
            "CONSULT": sum(p.text.count("CONSULT") for p in doc.paragraphs),
            "RISK": sum(p.text.count("RISK") for p in doc.paragraphs),
            "dependency_phrase": sum(
                ("依赖个案" in p.text or "个案创建成功" in p.text) for p in doc.paragraphs
            ),
        },
    )
    with ZipFile(path) as archive:
        bad_xml = []
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            value = archive.read(name).decode("utf-8", "ignore")
            if OLD.search(value):
                bad_xml.append(name)
        print("OLD_XML_PARTS", bad_xml)
